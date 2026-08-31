from django.test import TestCase, Client
from django.urls import reverse
from decimal import Decimal
from unittest.mock import patch
from django.contrib.auth.models import User
from panel.models import Empresa
from recursos_humanos.models import Nomina, Empleado, Contrato, Contratista, Beneficiario
from recursos_humanos.sat_service import SATService

class XMLPayrollParsingTest(TestCase):
    def setUp(self):
        self.empresa = Empresa.objects.create(
            nombre="Empresa de Prueba",
            subdominio="prueba",
            usuario_admin="admin",
            correo_contacto="prueba@test.com"
        )
        
    @patch('recursos_humanos.sat_service.CFDI')
    def test_parsear_y_guardar_xml_percepciones(self, mock_cfdi):
        # Configurar el mock de CFDI
        mock_cfdi.from_string.return_value = "mock_cfdi_instance"
        
        xml_content = """<cfdi:Comprobante xmlns:cfdi="http://www.sat.gob.mx/cfd/3" xmlns:tfd="http://www.sat.gob.mx/TimbreFiscalDigital" xmlns:nomina12="http://www.sat.gob.mx/nomina12" Version="3.3" Total="2000.00" Folio="999" Serie="N">
  <cfdi:Emisor Rfc="XAXX010101000"/>
  <cfdi:Receptor Rfc="POT120101XYZ" Nombre="PEDRO PEREZ"/>
  <cfdi:Complemento>
    <tfd:TimbreFiscalDigital UUID="e07b8b2e-1111-2222-3333-444455556666"/>
    <nomina12:Nomina FechaPago="2026-06-15" FechaInicialPago="2026-06-01" FechaFinalPago="2026-06-15" NumDiasPagados="15.00" TipoNomina="O">
      <nomina12:Receptor Curp="PEPE000000HDFRXX01" NumSeguridadSocial="12345678901"/>
      <nomina12:Percepciones>
        <nomina12:Percepcion TipoPercepcion="001" Clave="P001" Concepto="SUELDO QUINCENAL" ImporteGravado="1200.00" ImporteExento="0.00"/>
        <nomina12:Percepcion TipoPercepcion="002" Clave="P004" Concepto="AGUINALDO ANUAL" ImporteGravado="0.00" ImporteExento="500.00"/>
        <nomina12:Percepcion TipoPercepcion="001" Clave="P009" Concepto="PRIMA DE VACACIONES" ImporteGravado="250.00" ImporteExento="0.00"/>
        <nomina12:Percepcion TipoPercepcion="001" Clave="P009" Concepto="VACACIONES DIGNAS EX" ImporteGravado="0.00" ImporteExento="150.00"/>
      </nomina12:Percepciones>
    </nomina12:Nomina>
  </cfdi:Complemento>
</cfdi:Comprobante>"""
        
        class DummySATService(SATService):
            def __init__(self):
                pass
                
        service = DummySATService()
        
        # Ejecutar el parser
        resultado = service._parsear_y_guardar_xml(xml_content.encode('utf-8'), self.empresa, None)
        self.assertTrue(resultado)
        
        # Verificar que la nómina se guardó en la base de datos con desglose de percepciones
        nomina = Nomina.objects.get(uuid="e07b8b2e-1111-2222-3333-444455556666")
        
        self.assertEqual(nomina.nombre, "PEDRO PEREZ")
        self.assertEqual(nomina.rfc, "POT120101XYZ")
        self.assertEqual(nomina.curp, "PEPE000000HDFRXX01")
        self.assertEqual(nomina.nss, "12345678901")
        self.assertEqual(nomina.rfc_contratista, "XAXX010101000")
        self.assertEqual(nomina.periodo, "SAT Bimestre 3 2026")
        
        # Totales de percepciones desglosados
        self.assertEqual(nomina.sueldo_gravado, Decimal('1200.00'))
        self.assertEqual(nomina.aguinaldo_exento, Decimal('500.00'))
        self.assertEqual(nomina.aguinaldo_gravado, Decimal('0.00'))
        self.assertEqual(nomina.vacaciones_gravado, Decimal('250.00'))
        self.assertEqual(nomina.vacaciones_exento, Decimal('0.00'))
        self.assertEqual(nomina.vacaciones_dignas_gravado, Decimal('0.00'))
        self.assertEqual(nomina.vacaciones_dignas_exento, Decimal('150.00'))
        
        # Verificar JSON detallado y total_percepciones
        self.assertEqual(nomina.percepciones_detalladas.get('001', {}).get('gravado'), 1450.0)
        self.assertEqual(nomina.percepciones_detalladas.get('001', {}).get('exento'), 150.0)
        self.assertEqual(nomina.percepciones_detalladas.get('002', {}).get('exento'), 500.0)
        self.assertEqual(nomina.total_percepciones, Decimal('2100.00'))

    def test_crear_editar_nomina_percepciones_detalladas(self):
        # Iniciar sesión / crear datos mínimos
        from django.test import Client
        client = Client()
        client.force_login(User.objects.create_superuser(username="test_admin@prueba", email="test@test.com", password="pwd"))
        
        # Test creation via AJAX
        post_data = {
            'periodo': 'Bimestre 1 2026',
            'nombre': 'JUAN PEREZ',
            'rfc': 'XAXX010101000',
            'curp': 'XAXX010101HDFRXX01',
            'nss': '12345678901',
            'p_001_gravado': '1500.00',
            'p_001_exento': '0.00',
            'p_029_gravado': '0.00',
            'p_029_exento': '300.00'
        }
        
        # Simular sesión de sucursal
        session = client.session
        session['sucursal_id'] = None
        session.save()
        
        response = client.post('/recursos-humanos/nomina/crear/', post_data)
        self.assertEqual(response.status_code, 200)
        resp_json = response.json()
        self.assertTrue(resp_json['success'])
        
        # Verificar que se guardó correctamente
        nomina = Nomina.objects.get(nombre='JUAN PEREZ')
        self.assertEqual(nomina.sueldo_gravado, Decimal('1500.00')) # Sincronizado a legacy
        self.assertEqual(nomina.percepciones_detalladas.get('029', {}).get('exento'), 300.0)
        self.assertEqual(nomina.total_percepciones, Decimal('1800.00'))

class SISUBExportTest(TestCase):
    def setUp(self):
        self.empresa = Empresa.objects.create(
            nombre="Empresa de Prueba",
            subdominio="prueba",
            usuario_admin="admin",
            correo_contacto="prueba@test.com"
        )
        # Habilitar el módulo de Recursos Humanos para la empresa
        self.empresa.modulo_recursos_humanos = True
        self.empresa.save()
        
        self.user = User.objects.create_superuser(
            username="admin@prueba",
            email="admin@prueba.com",
            password="password"
        )
        self.contratista = Contratista.objects.create(
            empresa=self.empresa,
            rfc="CON010101AAA",
            nombre_razon_social="Contratista Prueba",
            correo="contratista@test.com"
        )
        self.empleado = Empleado.objects.create(
            empresa=self.empresa,
            num_empleado="E001",
            nss="12345678901",
            curp="PEPE000000HDFRXX01",
            rfc="PEPE000000XX1",
            nombre="Pedro",
            apellido_paterno="Perez",
            estado="activo"
        )
        # Crear un contrato con beneficiario = None
        self.contrato = Contrato.objects.create(
            empresa=self.empresa,
            contratista=self.contratista,
            beneficiario=None,
            fecha_inicio="2026-06-01",
            estado="vigente"
        )
        self.contrato.empleados.add(self.empleado)
        
        # Crear un recibo de nómina que caiga en el mes de Junio (Cuatrimestre 2)
        self.nomina = Nomina.objects.create(
            empresa=self.empresa,
            empleado=self.empleado,
            periodo="SAT Bimestre 3 2026",
            fecha_pago="2026-06-15",
            rfc="PEPE000000XX1",
            curp="PEPE000000HDFRXX01",
            nss="12345678901",
            nombre="Pedro Perez",
            sueldo_gravado=Decimal('1000.00'),
            dias_pagados=Decimal('15.00')
        )
        
    def test_exportar_sisub_sin_beneficiario(self):
        self.client.login(username="admin@prueba", password="password")
        from django.urls import reverse
        url = reverse('exportar_sisub_trabajadores', args=[self.contratista.id])
        response = self.client.get(url, {'cuatrimestre': 2, 'anio': 2026, 'formato': 'excel'})
        # Si tiene el error de ben.calle, responderá con error 500 (o lanzará una excepción en tests)
        self.assertEqual(response.status_code, 200)

    def test_exportar_sisub_incluye_contrato_no_vigente(self):
        # Cambiar el contrato a estado 'vencido'
        self.contrato.estado = 'vencido'
        self.contrato.save()

        self.client.login(username="admin@prueba", password="password")
        from django.urls import reverse
        url = reverse('exportar_sisub_trabajadores', args=[self.contratista.id])
        response = self.client.get(url, {'cuatrimestre': 2, 'anio': 2026, 'formato': 'csv'})
        self.assertEqual(response.status_code, 200)

        # Analizar el contenido del CSV
        content = response.content.decode('utf-8-sig')
        lines = [line.strip() for line in content.splitlines() if line.strip()]
        
        # Deben haber 2 líneas (cabecera + datos) porque el contrato 'vencido' fue incluido debido a las fechas
        self.assertEqual(len(lines), 2)
        row_data = lines[1].split(',')
        self.assertEqual(row_data[6], '12345678901') # NSS del trabajador

    def test_exportar_sisub_only_assigned_workers(self):
        # Crear un empleado que pertenece al contratista, pero no está asignado al contrato ManyToMany
        empleado_fallback = Empleado.objects.create(
            empresa=self.empresa,
            contratista=self.contratista,
            num_empleado="E002",
            nss="98765432101",
            curp="JUAN000000HDFRXX02",
            rfc="JUAN000000XX2",
            nombre="Juan",
            apellido_paterno="Gomez",
            estado="activo"
        )
        # Crear recibo de nómina para este empleado en Junio 2026 (Cuatrimestre 2)
        Nomina.objects.create(
            empresa=self.empresa,
            empleado=None,
            periodo="SAT Bimestre 3 2026",
            fecha_pago="2026-06-15",
            rfc="JUAN000000XX2",
            curp="JUAN000000HDFRXX02",
            nss="98765432101",
            nombre="Juan Gomez",
            sueldo_gravado=Decimal('800.00'),
            dias_pagados=Decimal('15.00')
        )

        self.client.login(username="admin@prueba", password="password")
        from django.urls import reverse
        url = reverse('exportar_sisub_trabajadores', args=[self.contratista.id])
        response = self.client.get(url, {'cuatrimestre': 2, 'anio': 2026, 'formato': 'csv'})
        self.assertEqual(response.status_code, 200)

        # Analizar el contenido del CSV
        content = response.content.decode('utf-8-sig')
        lines = [line.strip() for line in content.splitlines() if line.strip()]
        
        # No debe incluir a Juan Gomez porque no está asignado al contrato (solo Pedro Perez + cabecera = 2 líneas)
        self.assertEqual(len(lines), 2)
        
        # Verificar que el NSS de Juan Gomez (98765432101) NO esté presente
        nss_list = [line.split(',')[6] for line in lines[1:]]
        self.assertNotIn('98765432101', nss_list)

    def test_exportar_sisub_consolida_percepciones(self):
        # Crear segunda nomina en el mismo bimestre (Junio 2026 -> Bimestre 3)
        Nomina.objects.create(
            empresa=self.empresa,
            empleado=self.empleado,
            periodo="SAT Bimestre 3 2026",
            fecha_pago="2026-06-30",
            rfc="PEPE000000XX1",
            curp="PEPE000000HDFRXX01",
            nss="12345678901",
            nombre="Pedro Perez",
            sueldo_gravado=Decimal('1500.00'),
            dias_pagados=Decimal('15.00')
        )
        
        self.client.login(username="admin@prueba", password="password")
        from django.urls import reverse
        url = reverse('exportar_sisub_trabajadores', args=[self.contratista.id])
        response = self.client.get(url, {'cuatrimestre': 2, 'anio': 2026, 'formato': 'csv'})
        
        self.assertEqual(response.status_code, 200)
        
        # Analizar el contenido del CSV
        content = response.content.decode('utf-8-sig') # Decodificar omitiendo el BOM
        lines = [line.strip() for line in content.splitlines() if line.strip()]
        
        # Deben haber 2 líneas: 1 de cabecera y 1 de datos (por consolidación)
        self.assertEqual(len(lines), 2)
        
        # La línea de datos debe contener la suma de percepciones (1000 + 1500 = 2500)
        row_data = lines[1].split(',')
        self.assertEqual(row_data[15], '2500')

    def test_exportar_sisub_sdi_cero_fallback(self):
        from recursos_humanos.models import ImportacionSUA, TrabajadorSUA
        
        # Crear importación de SUA para Mayo 2026 (mes 5, Bimestre 3) con SDI > 0
        sua_mayo = ImportacionSUA.objects.create(
            empresa=self.empresa,
            periodo="MAYO 2026",
            tipo="mensual",
            rfc_empresa="CON010101AAA"
        )
        TrabajadorSUA.objects.create(
            importacion=sua_mayo,
            nss="12345678901",
            nombre="Pedro Perez",
            rfc_curp="PEPE000000HDFRXX01",
            sdi=Decimal('185.50')
        )
        
        # Crear importación de SUA para Junio 2026 (mes 6, Bimestre 3) con SDI = 0
        sua_junio = ImportacionSUA.objects.create(
            empresa=self.empresa,
            periodo="JUNIO 2026",
            tipo="mensual",
            rfc_empresa="CON010101AAA"
        )
        TrabajadorSUA.objects.create(
            importacion=sua_junio,
            nss="12345678901",
            nombre="Pedro Perez",
            rfc_curp="PEPE000000HDFRXX01",
            sdi=Decimal('0.00')
        )
        
        self.client.login(username="admin@prueba", password="password")
        from django.urls import reverse
        url = reverse('exportar_sisub_trabajadores', args=[self.contratista.id])
        response = self.client.get(url, {'cuatrimestre': 2, 'anio': 2026, 'formato': 'csv'})
        
        self.assertEqual(response.status_code, 200)
        
        content = response.content.decode('utf-8-sig')
        lines = [line.strip() for line in content.splitlines() if line.strip()]
        
        # La segunda línea (datos) debe tener el SDI de Mayo (185.50) en lugar del 0.00 de Junio
        row_data = lines[1].split(',')
        self.assertEqual(Decimal(row_data[18]), Decimal('185.50'))

    def test_exportar_sisub_suma_incapacidades(self):
        from recursos_humanos.models import ImportacionSUA, TrabajadorSUA
        
        # Crear importación de SUA para Mayo 2026 con 5 incapacidades
        sua_mayo = ImportacionSUA.objects.create(
            empresa=self.empresa,
            periodo="MAYO 2026",
            tipo="mensual",
            rfc_empresa="CON010101AAA"
        )
        TrabajadorSUA.objects.create(
            importacion=sua_mayo,
            nss="12-34-56-7890-1",  # Con guiones para probar normalización
            nombre="Pedro Perez",
            rfc_curp="PEPE000000HDFRXX01",
            incapacidades=5,
            sdi=Decimal('185.50')
        )
        
        # Crear importación de SUA para Junio 2026 con 3 incapacidades
        sua_junio = ImportacionSUA.objects.create(
            empresa=self.empresa,
            periodo="JUNIO 2026",
            tipo="mensual",
            rfc_empresa="CON010101AAA"
        )
        TrabajadorSUA.objects.create(
            importacion=sua_junio,
            nss="12345678901",
            nombre="Pedro Perez",
            rfc_curp="PEPE000000HDFRXX01",
            incapacidades=3,
            sdi=Decimal('185.50')
        )
        
        self.client.login(username="admin@prueba", password="password")
        from django.urls import reverse
        url = reverse('exportar_sisub_trabajadores', args=[self.contratista.id])
        response = self.client.get(url, {'cuatrimestre': 2, 'anio': 2026, 'formato': 'csv'})
        
        self.assertEqual(response.status_code, 200)
        
        content = response.content.decode('utf-8-sig')
        lines = [line.strip() for line in content.splitlines() if line.strip()]
        
        row_data = lines[1].split(',')
        # Columna 16 (0-indexed) de incapacidades debe ser la suma (5 + 3 = 8)
        self.assertEqual(int(row_data[16]), 8)

    def test_exportar_sisub_percepciones_fijas_variables(self):
        # Crear un recibo de nómina con percepciones detalladas en Junio 2026 (Cuatrimestre 2)
        # Percepciones fijas: '001' (gravado 1000, exento 100) -> Total 1100
        # Percepciones variables: '010' (gravado 500), '019' (exento 200) -> Total 700
        self.nomina.percepciones_detalladas = {
            '001': {'gravado': 1000.0, 'exento': 100.0},
            '010': {'gravado': 500.0, 'exento': 0.0},
            '019': {'gravado': 0.0, 'exento': 200.0}
        }
        self.nomina.save()
        
        self.client.login(username="admin@prueba", password="password")
        from django.urls import reverse
        url = reverse('exportar_sisub_trabajadores', args=[self.contratista.id])
        response = self.client.get(url, {'cuatrimestre': 2, 'anio': 2026, 'formato': 'csv'})
        
        self.assertEqual(response.status_code, 200)
        
        content = response.content.decode('utf-8-sig')
        lines = [line.strip() for line in content.splitlines() if line.strip()]
        
        row_data = lines[1].split(',')
        # Columna 14 (variables): 500 + 200 = 700
        self.assertEqual(row_data[14], '700')
        # Columna 15 (fijas): 1000 + 100 = 1100
        self.assertEqual(row_data[15], '1100')

    def test_exportar_sisub_percepciones_variables_switch(self):
        self.nomina.percepciones_detalladas = {
            '010': {'gravado': 500.0, 'exento': 0.0},
            '019': {'gravado': 0.0, 'exento': 200.0}
        }
        self.nomina.save()
        
        self.client.login(username="admin@prueba", password="password")
        from django.urls import reverse
        url = reverse('exportar_sisub_trabajadores', args=[self.contratista.id])
        
        # 1. Con switch apagado (percepciones_var = '0')
        response_off = self.client.get(url, {'cuatrimestre': 2, 'anio': 2026, 'formato': 'csv', 'percepciones_var': '0'})
        self.assertEqual(response_off.status_code, 200)
        lines_off = [l.strip() for l in response_off.content.decode('utf-8-sig').splitlines() if l.strip()]
        row_off = lines_off[1].split(',')
        self.assertEqual(row_off[14], '0')  # Debe mostrar 0 en columna O
        
        # 2. Con switch encendido (percepciones_var = '1')
        response_on = self.client.get(url, {'cuatrimestre': 2, 'anio': 2026, 'formato': 'csv', 'percepciones_var': '1'})
        self.assertEqual(response_on.status_code, 200)
        lines_on = [l.strip() for l in response_on.content.decode('utf-8-sig').splitlines() if l.strip()]
        row_on = lines_on[1].split(',')
        self.assertEqual(row_on[14], '700') # Debe mostrar la suma real

    def test_exportar_sisub_vales_despensa_uma_split(self):
        self.empresa.uma = Decimal('117.31')
        self.empresa.save()
        
        import datetime
        self.nomina.fecha_pago = datetime.date(2026, 5, 15)  # Mayo tiene 31 días
        self.nomina.dias_pagados = Decimal('31.00')
        self.nomina.percepciones_detalladas = {
            '029': {'gravado': 5000.00, 'exento': 0.0}
        }
        self.nomina.save()
        
        self.client.login(username="admin@prueba", password="password")
        from django.urls import reverse
        url = reverse('exportar_sisub_trabajadores', args=[self.contratista.id])
        response = self.client.get(url, {'cuatrimestre': 2, 'anio': 2026, 'formato': 'csv'})
        
        self.assertEqual(response.status_code, 200)
        
        content = response.content.decode('utf-8-sig')
        lines = [line.strip() for line in content.splitlines() if line.strip()]
        
        row_data = lines[1].split(',')
        # Columna 14 (variables): 5000 - (0.40 * 117.31 * 31) = 3545.36 -> rounded = 3545
        self.assertEqual(row_data[14], '3545')
        # Columna 15 (fijas): 0 ya que los vales no integran a fijas en ningún caso
        self.assertEqual(row_data[15], '0')
        # Columna 17 (no integrables): 0.40 * 117.31 * 31 = 1454.64 -> rounded = 1455
        self.assertEqual(row_data[17], '1455')

    def test_exportar_sisub_conceptos_especiales(self):
        import datetime
        from django.urls import reverse
        self.client.login(username="admin@prueba", password="password")
        url = reverse('exportar_sisub_trabajadores', args=[self.contratista.id])

        # --- CASO 1: PTU a tiempo (Mayo 2026, Bimestre 3) ---
        self.nomina.fecha_pago = datetime.date(2026, 5, 15)
        self.nomina.fecha_inicial_pago = datetime.date(2026, 5, 1)
        self.nomina.fecha_final_pago = datetime.date(2026, 5, 15)
        self.nomina.dias_pagados = 15
        self.nomina.sdi = 200.00  # Límite PTU = 90 * 200 = 18000
        self.nomina.percepciones_detalladas = {
            '003': {'gravado': 25000.00, 'exento': 0.0}
        }
        self.nomina.deducciones_detalladas = {}
        self.nomina.save()
        
        response = self.client.get(url, {'cuatrimestre': 2, 'anio': 2026, 'formato': 'csv'})
        self.assertEqual(response.status_code, 200)
        row = response.content.decode('utf-8-sig').splitlines()[1].split(',')
        self.assertEqual(row[14], '7000')  # var: 25000 - 18000 = 7000
        self.assertEqual(row[17], '18000') # no_int: 18000

        # --- CASO 2: PTU fuera de plazo (Junio 2026) ---
        self.nomina.fecha_pago = datetime.date(2026, 6, 15)
        self.nomina.fecha_inicial_pago = datetime.date(2026, 6, 1)
        self.nomina.fecha_final_pago = datetime.date(2026, 6, 15)
        self.nomina.percepciones_detalladas = {
            '003': {'gravado': 25000.00, 'exento': 0.0}
        }
        self.nomina.save()
        
        response = self.client.get(url, {'cuatrimestre': 2, 'anio': 2026, 'formato': 'csv'})
        row = response.content.decode('utf-8-sig').splitlines()[1].split(',')
        self.assertEqual(row[14], '25000') # var: todo integra por fuera de plazo
        self.assertEqual(row[17], '0')     # no_int: 0

        # --- CASO 3: Fondo de Ahorro sin excedentes (Junio 2026) ---
        self.nomina.fecha_pago = datetime.date(2026, 6, 15)
        self.nomina.fecha_inicial_pago = datetime.date(2026, 6, 1)
        self.nomina.fecha_final_pago = datetime.date(2026, 6, 15)
        self.nomina.percepciones_detalladas = {
            '005': {'gravado': 0.0, 'exento': 1000.00}
        }
        self.nomina.deducciones_detalladas = {
            '005': {'importe': 1000.00}
        }
        self.nomina.retiros_ahorro_excedido = False
        self.nomina.save()
        
        response = self.client.get(url, {'cuatrimestre': 2, 'anio': 2026, 'formato': 'csv'})
        row = response.content.decode('utf-8-sig').splitlines()[1].split(',')
        self.assertEqual(row[14], '0')     # var: 0
        self.assertEqual(row[17], '1000')  # no_int: 1000

        # --- CASO 4: Fondo de Ahorro con excedentes patronal (Junio 2026) ---
        self.nomina.percepciones_detalladas = {
            '005': {'gravado': 0.0, 'exento': 1500.00}
        }
        self.nomina.deducciones_detalladas = {
            '005': {'importe': 1000.00}
        }
        self.nomina.save()
        
        response = self.client.get(url, {'cuatrimestre': 2, 'anio': 2026, 'formato': 'csv'})
        row = response.content.decode('utf-8-sig').splitlines()[1].split(',')
        self.assertEqual(row[14], '500')    # var: 1500 - 1000 = 500
        self.assertEqual(row[17], '1000')   # no_int: 1000

        # --- CASO 5: Fondo de Ahorro con retiros excedidos (Junio 2026) ---
        self.nomina.retiros_ahorro_excedido = True
        self.nomina.save()
        
        response = self.client.get(url, {'cuatrimestre': 2, 'anio': 2026, 'formato': 'csv'})
        row = response.content.decode('utf-8-sig').splitlines()[1].split(',')
        self.assertEqual(row[14], '1500')   # var: todo integra
        self.assertEqual(row[17], '0')      # no_int: 0

        # --- CASO 6: Alimentación onerosa (Descuento >= 20% UMA) (Junio 2026) ---
        # UMA = 117.31. Dias = 15. Min discount = 0.20 * 117.31 * 15 = 351.93
        self.nomina.retiros_ahorro_excedido = False
        self.nomina.dias_pagados = 15
        self.nomina.percepciones_detalladas = {
            '047': {'gravado': 3000.00, 'exento': 0.0}
        }
        self.nomina.deducciones_detalladas = {
            '007': {'importe': 360.00}  # 360 >= 351.93
        }
        self.nomina.save()
        
        response = self.client.get(url, {'cuatrimestre': 2, 'anio': 2026, 'formato': 'csv'})
        row = response.content.decode('utf-8-sig').splitlines()[1].split(',')
        self.assertEqual(row[15], '0')     # fijas: 0
        self.assertEqual(row[17], '3000')  # no_int: 3000

        # --- CASO 7: Alimentación no onerosa (Descuento < 20% UMA) (Junio 2026) ---
        self.nomina.deducciones_detalladas = {
            '007': {'importe': 300.00}  # 300 < 351.93
        }
        self.nomina.save()
        
        response = self.client.get(url, {'cuatrimestre': 2, 'anio': 2026, 'formato': 'csv'})
        row = response.content.decode('utf-8-sig').splitlines()[1].split(',')
        self.assertEqual(row[15], '3000')  # fijas: 3000 (integra completo)
        self.assertEqual(row[17], '0')     # no_int: 0

    def test_exportar_sisub_excluye_cancelados(self):
        import datetime
        from django.urls import reverse
        self.client.login(username="admin@prueba", password="password")
        url = reverse('exportar_sisub_trabajadores', args=[self.contratista.id])

        # Crear una nómina vigente para el empleado
        Nomina.objects.create(
            empresa=self.empresa,
            empleado=self.empleado,
            periodo="Mayo 2026",
            tipo_nomina="O",
            fecha_pago=datetime.date(2026, 5, 15),
            fecha_inicial_pago=datetime.date(2026, 5, 1),
            fecha_final_pago=datetime.date(2026, 5, 15),
            dias_pagados=15,
            sdi=150.00,
            nss=self.empleado.nss,
            curp=self.empleado.curp,
            nombre=f"{self.empleado.nombre} {self.empleado.apellido_paterno} {self.empleado.apellido_materno}".upper(),
            sueldo_gravado=Decimal("5000.00"),
            estado="vigente"
        )

        # Crear una nómina CANCELADA para el empleado en el mismo periodo
        Nomina.objects.create(
            empresa=self.empresa,
            empleado=self.empleado,
            periodo="Mayo 2026",
            tipo_nomina="O",
            fecha_pago=datetime.date(2026, 5, 15),
            fecha_inicial_pago=datetime.date(2026, 5, 1),
            fecha_final_pago=datetime.date(2026, 5, 15),
            dias_pagados=15,
            sdi=150.00,
            nss=self.empleado.nss,
            curp=self.empleado.curp,
            nombre=f"{self.empleado.nombre} {self.empleado.apellido_paterno} {self.empleado.apellido_materno}".upper(),
            sueldo_gravado=Decimal("99999.00"), # Importe muy alto para detectar si se suma por error
            estado="cancelado"
        )

        response = self.client.get(url, {'cuatrimestre': 2, 'anio': 2026, 'formato': 'csv'})
        self.assertEqual(response.status_code, 200)
        
        content_str = response.content.decode('utf-8-sig')
        # Verificar que el sueldo gravado del cancelado (99999) NO esté en el reporte
        self.assertNotIn("99999", content_str)

    def test_exportar_sisub_contratos(self):
        import datetime
        # Configurar primer contrato
        self.contrato.folio = 'CON-002'
        self.contrato.fecha_inicio = datetime.date(2026, 6, 15)
        self.contrato.fecha_fin = datetime.date(2026, 8, 30)
        self.contrato.vigencia_contrato = datetime.date(2026, 8, 31)
        self.contrato.save()
        
        # Crear segundo contrato (debe ir primero en ordenamiento por folio)
        contrato_2 = Contrato.objects.create(
            empresa=self.empresa,
            contratista=self.contratista,
            beneficiario=None,
            folio='CON-001',
            tipo_contrato='01',
            fecha_inicio=datetime.date(2026, 5, 10),
            fecha_fin=datetime.date(2026, 6, 20),
            vigencia_contrato=datetime.date(2026, 7, 1),
            estado="vigente"
        )
        
        self.client.login(username="admin@prueba", password="password")
        from django.urls import reverse
        url = reverse('exportar_sisub_contratos', args=[self.contratista.id])
        response = self.client.get(url, {'cuatrimestre': 2, 'anio': 2026, 'formato': 'csv'})
        
        self.assertEqual(response.status_code, 200)
        
        # Analizar el contenido del CSV
        content = response.content.decode('utf-8-sig')
        lines = [line.strip() for line in content.splitlines() if line.strip()]
        
        self.assertEqual(len(lines), 3)
        
        row_1 = lines[1].split(',')
        row_2 = lines[2].split(',')
        
        # Verificar ordenamiento por folio
        self.assertEqual(row_1[3], 'CON-001')
        self.assertEqual(row_2[3], 'CON-002')
        
        # Verificar que el tipo de contrato para el código 01 haya perdido el prefijo
        self.assertEqual(row_1[4], 'Contrato de trabajo por tiempo indeterminado')
        
        # Verificar formato de fecha dd/mm/aaaa en Vigencia, Inicio, Término
        self.assertEqual(row_1[7], '01/07/2026')
        self.assertEqual(row_1[8], '10/05/2026')
        self.assertEqual(row_1[9], '20/06/2026')

    def test_contrato_save_string_dates(self):
        import datetime
        con = Contrato.objects.create(
            empresa=self.empresa,
            contratista=self.contratista,
            beneficiario=None,
            folio="CON-TEST-STR",
            tipo_contrato="01",
            fecha_inicio="2026-06-01",
            fecha_fin="2026-12-31",
            vigencia_contrato="2027-01-01",
            estado="vigente"
        )
        self.assertIsInstance(con.fecha_inicio, datetime.date)
        self.assertEqual(con.fecha_inicio, datetime.date(2026, 6, 1))
        
        self.assertIsInstance(con.fecha_fin, datetime.date)
        self.assertEqual(con.fecha_fin, datetime.date(2026, 12, 31))
        
        self.assertIsInstance(con.vigencia_contrato, datetime.date)
        self.assertEqual(con.vigencia_contrato, datetime.date(2027, 1, 1))

    def test_exportar_nominas_excel(self):
        # Crear un recibo con percepciones para exportar
        Nomina.objects.create(
            empresa=self.empresa,
            empleado=self.empleado,
            periodo="SAT Bimestre 3 2026",
            fecha_pago="2026-06-15",
            rfc="PEPE000000XX1",
            curp="PEPE000000HDFRXX01",
            nss="12345678901",
            nombre="Pedro Perez",
            sueldo_gravado=Decimal('1200.00'),
            vacaciones_gravado=Decimal('300.00'),
            vacaciones_exento=Decimal('100.00'),
            vacaciones_dignas_gravado=Decimal('400.00'),
            vacaciones_dignas_exento=Decimal('200.00'),
            aguinaldo_gravado=Decimal('500.00'),
            aguinaldo_exento=Decimal('250.00')
        )
        self.client.login(username="admin@prueba", password="password")
        from django.urls import reverse
        url = reverse('exportar_nominas_excel')
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        
        # Leer el Excel generado en memoria
        import io
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(response.content), data_only=True)
        ws = wb.active
        
        # Verificar cabeceras
        headers = [cell.value for cell in ws[1]]
        self.assertIn("Sueldo (Gravado)", headers)
        self.assertIn("Vacaciones (Gravado)", headers)
        self.assertIn("Vacaciones (Exento)", headers)
        self.assertIn("Vacaciones Dignas (Gravado)", headers)
        self.assertIn("Vacaciones Dignas (Exento)", headers)
        self.assertIn("Aguinaldo (Gravado)", headers)
        self.assertIn("Aguinaldo (Exento)", headers)

class ContratoEstadosYVersionesTest(TestCase):
    def setUp(self):
        self.empresa = Empresa.objects.create(
            nombre="Empresa de Prueba",
            subdominio="prueba",
            usuario_admin="admin",
            correo_contacto="prueba@test.com"
        )
        self.contratista = Contratista.objects.create(
            empresa=self.empresa,
            rfc="CON010101AAA",
            nombre_razon_social="Contratista Prueba",
            correo="contratista@test.com"
        )

    def test_contrato_estados_y_versiones(self):
        import datetime
        today = datetime.date.today()
        past_date = today - datetime.timedelta(days=10)
        future_date = today + datetime.timedelta(days=10)

        # 1. Contrato vigente en fechas y periodicidad
        contrato_vigente = Contrato.objects.create(
            empresa=self.empresa,
            contratista=self.contratista,
            fecha_inicio=past_date,
            fecha_fin=future_date,
            version="1",
            estado_periodicidad="vigente"
        )
        self.assertEqual(contrato_vigente.estado_vigencia, 'vigente')
        self.assertEqual(contrato_vigente.estado, 'vigente')
        self.assertEqual(contrato_vigente.version, '1')

        # 2. Contrato con periodicidad cerrado (se auto-calcula por fecha_fin en el pasado)
        contrato_cerrado = Contrato.objects.create(
            empresa=self.empresa,
            contratista=self.contratista,
            fecha_inicio=past_date - datetime.timedelta(days=5),
            fecha_fin=past_date,
            version="2",
            estado_periodicidad="vigente"
        )
        self.assertEqual(contrato_cerrado.estado_periodicidad, 'cerrado')
        self.assertEqual(contrato_cerrado.estado, 'cerrado')
        self.assertEqual(contrato_cerrado.version, '2')

        # 3. Contrato vencido en fechas (vigencia_contrato y fecha_fin en el pasado)
        contrato_vencido = Contrato.objects.create(
            empresa=self.empresa,
            contratista=self.contratista,
            fecha_inicio=past_date - datetime.timedelta(days=10),
            fecha_fin=past_date,
            vigencia_contrato=past_date,
            version="1",
            estado_periodicidad="vigente"
        )
        self.assertEqual(contrato_vencido.estado_vigencia, 'vencido')
        self.assertEqual(contrato_vencido.estado_periodicidad, 'cerrado')
        self.assertEqual(contrato_vencido.estado, 'vencido')

    def test_contrato_versionamiento_por_contratista(self):
        import datetime
        today = datetime.date.today()
        future_date = today + datetime.timedelta(days=10)

        # Crear un contratista B
        contratista_b = Contratista.objects.create(
            empresa=self.empresa,
            rfc="CON020202BBB",
            nombre_razon_social="Contratista B",
            correo="contratistab@test.com"
        )

        # 1. Contrato Contratista A, Folio 'FOL-123'
        c_a1 = Contrato.objects.create(
            empresa=self.empresa,
            contratista=self.contratista,
            folio="FOL-123",
            fecha_inicio=today,
            fecha_fin=future_date,
            vigencia_contrato=future_date
        )

        # 2. Contrato Contratista B, Folio 'FOL-123' (Mismo folio, diferente contratista)
        c_b1 = Contrato.objects.create(
            empresa=self.empresa,
            contratista=contratista_b,
            folio="FOL-123",
            fecha_inicio=today,
            fecha_fin=future_date,
            vigencia_contrato=future_date
        )

        # Ambos deben ser Version 1 y vigentes
        c_a1.refresh_from_db()
        c_b1.refresh_from_db()
        self.assertEqual(c_a1.version, '1')
        self.assertEqual(c_a1.estado_vigencia, 'vigente')
        self.assertEqual(c_b1.version, '1')
        self.assertEqual(c_b1.estado_vigencia, 'vigente')

        # 3. Crear otra versión para Contratista A, Folio 'FOL-123'
        c_a2 = Contrato.objects.create(
            empresa=self.empresa,
            contratista=self.contratista,
            folio="FOL-123",
            fecha_inicio=future_date + datetime.timedelta(days=1),
            fecha_fin=future_date + datetime.timedelta(days=10),
            vigencia_contrato=future_date + datetime.timedelta(days=10)
        )

        c_a1.refresh_from_db()
        c_a2.refresh_from_db()
        c_b1.refresh_from_db()

        # Contratista A: c_a1 debe cerrarse (version 1), c_a2 debe ser vigente (version 2)
        self.assertEqual(c_a1.version, '1')
        self.assertEqual(c_a1.estado_vigencia, 'cerrado')
        self.assertEqual(c_a2.version, '2')
        self.assertEqual(c_a2.estado_vigencia, 'vigente')

        # Contratista B: c_b1 no debe verse afectado (sigue vigente y version 1)
        self.assertEqual(c_b1.version, '1')
        self.assertEqual(c_b1.estado_vigencia, 'vigente')


class SUAExportTest(TestCase):
    def setUp(self):
        self.empresa = Empresa.objects.create(
            nombre="Empresa de Prueba",
            subdominio="prueba",
            usuario_admin="admin",
            correo_contacto="prueba@test.com"
        )
        self.empresa.modulo_recursos_humanos = True
        self.empresa.save()
        
        self.user = User.objects.create_superuser(
            username="admin@prueba",
            email="admin@prueba.com",
            password="password"
        )
        
        from recursos_humanos.models import ImportacionSUA, TrabajadorSUA
        self.importacion = ImportacionSUA.objects.create(
            empresa=self.empresa,
            periodo="MAYO 2026",
            tipo="mensual",
            rfc_empresa="CON010101AAA",
            nombre_razon_social="Contratista A",
            registro_patronal="A1234567890"
        )
        TrabajadorSUA.objects.create(
            importacion=self.importacion,
            nss="12345678901",
            nombre="Pedro Perez",
            rfc_curp="PEPE000000HDFRXX01",
            sdi=Decimal('185.50')
        )
        
    def test_exportar_sua_excel_xlsx(self):
        self.client.login(username="admin@prueba", password="password")
        from django.urls import reverse
        url = reverse('exportar_sua_excel', args=[self.importacion.id])
        response = self.client.get(url, {'formato': 'excel'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    def test_exportar_sua_excel_csv(self):
        self.client.login(username="admin@prueba", password="password")
        from django.urls import reverse
        url = reverse('exportar_sua_excel', args=[self.importacion.id])
        response = self.client.get(url, {'formato': 'csv'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'text/csv')
        content = response.content.decode('utf-8-sig')
        self.assertIn("REPORTE DE INTEGRACIÓN SUA", content)
        self.assertIn("Pedro Perez", content)


class HRNotificationTest(TestCase):
    def setUp(self):
        from django.contrib.auth.models import User
        from panel.models import Empresa
        
        self.empresa = Empresa.objects.create(
            nombre="Empresa de Prueba",
            subdominio="prueba",
            usuario_admin="admin",
            correo_contacto="prueba@test.com"
        )
        self.empresa.modulo_recursos_humanos = True
        self.empresa.save()
        
        self.admin = User.objects.create_superuser(
            username="admin@prueba",
            email="admin@prueba.com",
            password="password"
        )
        self.normal_user = User.objects.create_user(
            username="user1@prueba",
            email="user1@prueba.com",
            password="password"
        )
        
    def test_crear_empleado_generates_notification(self):
        from notificaciones.models import Notificacion
        self.client.login(username="admin@prueba", password="password")
        from django.urls import reverse
        url = reverse('crear_empleado_ajax')
        
        response = self.client.post(url, {
            'num_empleado': 'EMP-001',
            'nombre': 'Juan',
            'apellido_paterno': 'Perez',
            'apellido_materno': 'Gomez',
            'curp': 'CURP000000HDFRXX01',
            'rfc': 'RFC0000000XXX',
            'nss': '12345678901',
            'salario_diario_ordinario': '200.00',
            'sbc': '200.00',
            'sdi': '200.00',
            'forma_pago': 'quincenal',
            'tipo_salario': 'fijo',
            'puesto': 'Puesto de Prueba',
            'departamento': 'Dept de Prueba',
            'estado': 'activo',
            'clave_ubicacion': 'BEN01',
            'genero': 'H',
            'estado_civil': 'soltero',
            'correo_personal': 'test@test.com',
            'telefono_movil': '1234567890',
            'nacionalidad': 'Mexicana',
            'riesgo_trabajo': 'I',
            'tipo_trabajador': 'base',
            'tipo_contrato': '01',
            'jornada': 'diurna',
            'unidad_monetaria': 'MXN',
            'clave_percepcion_sat': '001',
            'tipo_cuenta': 'tarjeta'
        })
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()['success'])
        
        notifs = Notificacion.objects.filter(empresa=self.empresa)
        self.assertEqual(notifs.count(), 1)
        notif = notifs.first()
        self.assertEqual(notif.actor, self.admin)
        self.assertEqual(notif.propietario_recurso, self.admin)
        self.assertIn("creó al empleado Juan Perez", notif.mensaje)


class ICSOEExportTest(TestCase):
    def setUp(self):
        from django.contrib.auth.models import User
        from panel.models import Empresa
        from recursos_humanos.models import Contratista, Beneficiario, Contrato, ImportacionSUA, TrabajadorSUA
        
        self.empresa = Empresa.objects.create(
            nombre="Empresa de Prueba",
            subdominio="prueba",
            usuario_admin="admin",
            correo_contacto="prueba@test.com"
        )
        self.empresa.modulo_recursos_humanos = True
        self.empresa.save()
        
        self.admin = User.objects.create_superuser(
            username="admin@prueba",
            email="admin@prueba.com",
            password="password"
        )
        
        # Crear contratista
        self.contratista = Contratista.objects.create(
            empresa=self.empresa,
            rfc="CON010101AAA",
            nombre_razon_social="Contratista A",
            correo="contratista@test.com",
            registro_patronal="A1234567890"
        )
        
        # Crear beneficiarios
        self.beneficiario1 = Beneficiario.objects.create(
            empresa=self.empresa,
            rfc="BEN010101AAA",
            nombre_razon_social="Beneficiario 1",
            clave="BEN01",
            correo="b1@test.com"
        )
        self.beneficiario2 = Beneficiario.objects.create(
            empresa=self.empresa,
            rfc="BEN020202BBB",
            nombre_razon_social="Beneficiario 2",
            clave="BEN02",
            correo="b2@test.com"
        )
        
        # Crear contrato vinculado a beneficiario1
        import datetime
        self.contrato = Contrato.objects.create(
            empresa=self.empresa,
            contratista=self.contratista,
            beneficiario=self.beneficiario1,
            folio="CONTRATO-001",
            fecha_inicio=datetime.date(2026, 1, 1),
            fecha_fin=datetime.date(2026, 12, 31),
            vigencia_contrato=datetime.date(2026, 12, 31)
        )
        
        # Crear importación de SUA bimestral
        self.importacion = ImportacionSUA.objects.create(
            empresa=self.empresa,
            periodo="ABRIL 2026",
            tipo="bimestral",
            rfc_empresa="CON010101AAA",
            nombre_razon_social="Contratista A",
            registro_patronal="A1234567890"
        )
        
        # Trabajador 1 (tiene beneficiario del contrato BEN01, no tiene crédito)
        TrabajadorSUA.objects.create(
            importacion=self.importacion,
            nss="11111111111",
            nombre="Juan Perez",
            clave_ubicacion="BEN01",
            aportacion_patronal=Decimal('500.00'),
            tipo_valor_infonavit='-',
            amortizacion=Decimal('0.00')
        )
        # Trabajador 2 (tiene beneficiario del contrato BEN01, tiene crédito)
        TrabajadorSUA.objects.create(
            importacion=self.importacion,
            nss="22222222222",
            nombre="Maria Gomez",
            clave_ubicacion="BEN01",
            aportacion_patronal=Decimal('600.00'),
            tipo_valor_infonavit='FD',
            amortizacion=Decimal('150.00')
        )
        # Trabajador 3 (tiene beneficiario de otro contrato BEN02, no vinculado)
        TrabajadorSUA.objects.create(
            importacion=self.importacion,
            nss="33333333333",
            nombre="Pedro Ruiz",
            clave_ubicacion="BEN02",
            aportacion_patronal=Decimal('800.00'),
            tipo_valor_infonavit='-',
            amortizacion=Decimal('0.00')
        )

    def test_exportar_icsoe_calculations(self):
        self.client.login(username="admin@prueba", password="password")
        from django.urls import reverse
        url = reverse('exportar_icsoe', args=[self.contratista.id])
        
        # Cuatrimestre 1 (Ene-Abr 2026), formato CSV
        response = self.client.get(url, {'cuatrimestre': '1', 'anio': '2026', 'formato': 'csv'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'text/csv')
        
        # Decodificar contenido
        content = response.content.decode('utf-8-sig')
        # Verificar que solo sume los trabajadores con clave BEN01 (Juan y Maria)
        self.assertIn("500", content)
        self.assertIn("600", content)
        self.assertIn("150", content)
        # No debe haber 800 de Pedro
        self.assertNotIn("800", content)

class ImportarContratistasTest(TestCase):
    def setUp(self):
        from django.contrib.auth.models import User
        from panel.models import Empresa
        
        self.empresa = Empresa.objects.create(
            nombre="Empresa Importadora",
            subdominio="importadora",
            modulo_recursos_humanos=True
        )
        self.user = User.objects.create_superuser(username="admin@importadora", password="password")
        self.client = Client()
        self.client.force_login(self.user)
        
        session = self.client.session
        session['empresa_id'] = self.empresa.id
        session.save()

    def test_importar_contratistas_template_con_titulos(self):
        import openpyxl
        import io
        from django.urls import reverse
        from recursos_humanos.models import Contratista
        
        # Crear un archivo Excel en memoria simulando el formato generado con títulos arriba (los encabezados reales en la fila 3)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sujeto Obligado"
        
        # Fila 1: Título agrupador
        ws.append(["a-Datos Generales", "", "", "", "", ""])
        # Fila 2: Sub-agrupador
        ws.append(["Periodo", "b-Datos de identificacion", "", "", "", ""])
        # Fila 3: Encabezados reales
        headers = [
            "cuatrimestre que declara", 
            "anio que se declara", 
            "Registro Federal de Contribuyente", 
            "Nombre denominacion o razon social", 
            "Correo electronico", 
            "Registro patronal"
        ]
        ws.append(headers)
        # Fila 4: Datos
        data = [
            "1", 
            "2026", 
            "CON990909XYZ", 
            "Contratista Importado S.A.", 
            "importado@test.com", 
            "A999999999"
        ]
        ws.append(data)
        
        # Guardar en memoria
        excel_file = io.BytesIO()
        wb.save(excel_file)
        excel_file.seek(0)
        
        # Realizar la petición POST
        url = reverse('importar_contratistas_ajax')
        response = self.client.post(url, {'archivo': excel_file}, format='multipart')
        self.assertEqual(response.status_code, 200)
        
        resp_json = response.json()
        self.assertEqual(resp_json['status'], 'success')
        self.assertIn("Se registraron/actualizaron 1 contratistas", resp_json['message'])
        
        # Verificar en base de datos
        contratista = Contratista.objects.get(rfc="CON990909XYZ")
        self.assertEqual(contratista.nombre_razon_social, "Contratista Importado S.A.")
        self.assertEqual(contratista.correo, "importado@test.com")
        self.assertEqual(contratista.registro_patronal, "A999999999")

        # Verificar notificación
        from notificaciones.models import Notificacion
        notifs = Notificacion.objects.filter(empresa=self.empresa)
        self.assertEqual(notifs.count(), 1)
        notif = notifs.first()
        self.assertEqual(notif.actor, self.user)
        self.assertIn("importó de forma masiva 1 contratistas (nuevos: 1, actualizados: 0)", notif.mensaje)

    def test_importar_contratos_generates_notification(self):
        import openpyxl
        import io
        from django.urls import reverse
        from decimal import Decimal
        from recursos_humanos.models import Contrato, Beneficiario
        from notificaciones.models import Notificacion
        
        # Crear un archivo Excel en memoria para contratos
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Contratos"
        
        headers = [
            "Registro Federal de Contribuyentes",  # RFC Beneficiario
            "Registro Federal de Contribuyente del sujeto obligado",  # RFC Contratista
            "Objeto del contrato",
            "Nombre denominacion o razon social",  # Nombre Beneficiario
            "Registro Patronal ante el IMSS",
            "Calle", "Numero exterior", "Numero interior", "Entre calle", "Y calle", "Colonia",
            "Codigo Postal", "Municipio o Alcaldia", "Entidad Federativa", "Correo electronico",
            "telefono (numero extension)", "Numero de contrato", "Tipo de contrato", "Monto del contrato",
            "Vigencia (del contrato)", "Fecha de inicio (del contrato)", "Fecha de termino (del contrato)",
            "Numero estimado mensual de trabajadores que se pondran a disposicion (del contrato)"
        ]
        ws.append(headers)
        
        # Registro 1
        data = [
            "BEN111111AAA",
            "CON222222BBB",
            "Servicios de Limpieza Especializada",
            "Beneficiario Limpieza S.A.",
            "B12345678",
            "Calle Falsa", "123", "", "Entre 1", "Y 2", "Centro",
            "06000", "Cuauhtemoc", "CDMX", "limpieza@test.com",
            "5555555555", "CON-LIM-01", "indeterminado", "50000.00",
            "2026-12-31", "2026-01-01", "2026-12-31", "5"
        ]
        ws.append(data)
        
        # Guardar en memoria
        excel_file = io.BytesIO()
        wb.save(excel_file)
        excel_file.seek(0)
        
        # Realizar la petición POST
        url = reverse('importar_contratos_ajax')
        response = self.client.post(url, {'archivo': excel_file}, format='multipart')
        self.assertEqual(response.status_code, 200)
        
        resp_json = response.json()
        self.assertEqual(resp_json['status'], 'success')
        self.assertIn("Se registraron 1 contratos", resp_json['message'])
        
        # Verificar en base de datos
        contrato = Contrato.objects.get(folio="CON-LIM-01")
        self.assertEqual(contrato.monto_contrato, Decimal("50000.00"))
        
        beneficiario = Beneficiario.objects.get(rfc="BEN111111AAA")
        self.assertEqual(beneficiario.nombre_razon_social, "Beneficiario Limpieza S.A.")
        
        # Verificar notificación
        notifs = Notificacion.objects.filter(empresa=self.empresa)
        self.assertEqual(notifs.count(), 1)
        notif = notifs.first()
        self.assertEqual(notif.actor, self.user)
        self.assertIn("importó de forma masiva 1 contratos (nuevos: 1, actualizados: 0) y creó 1 beneficiarios", notif.mensaje)


class CargaTrabajadoresICSOETest(TestCase):
    def setUp(self):
        self.empresa = Empresa.objects.create(
            nombre="Empresa de Prueba",
            subdominio="prueba",
            usuario_admin="admin",
            correo_contacto="prueba@test.com"
        )
        self.empresa.modulo_recursos_humanos = True
        self.empresa.save()

        self.user = User.objects.create_superuser(
            username="admin@prueba",
            email="admin@prueba.com",
            password="password"
        )
        
        self.contratista = Contratista.objects.create(
            empresa=self.empresa,
            nombre_razon_social="Contratista S.A.",
            rfc="CON111111AAA"
        )
        
        self.beneficiario = Beneficiario.objects.create(
            empresa=self.empresa,
            nombre_razon_social="Beneficiario S.A.",
            rfc="BEN111111AAA"
        )
        
        import datetime
        self.contrato = Contrato.objects.create(
            empresa=self.empresa,
            contratista=self.contratista,
            beneficiario=self.beneficiario,
            folio="CON-TEST-100",
            fecha_inicio=datetime.date(2026, 1, 1),
            fecha_fin=datetime.date(2026, 4, 30),
            monto_contrato=10000.00
        )
        
        # Empleado 1: tiene el beneficiario directamente asignado en su perfil
        self.empleado_directo = Empleado.objects.create(
            empresa=self.empresa,
            nombre="Juan",
            apellido_paterno="Perez",
            nss="11111111111",
            curp="CURP11111111111111",
            contratista=self.contratista,
            beneficiario=self.beneficiario
        )
        
        # Empleado 2: no tiene el beneficiario asignado en su perfil, pero está asignado al contrato con el beneficiario
        self.empleado_contrato = Empleado.objects.create(
            empresa=self.empresa,
            nombre="Pedro",
            apellido_paterno="Gomez",
            nss="22222222222",
            curp="CURP22222222222222",
            contratista=self.contratista
        )
        self.contrato.empleados.add(self.empleado_contrato)
        
        # Crear importación SUA
        from recursos_humanos.models import ImportacionSUA, TrabajadorSUA
        self.importacion = ImportacionSUA.objects.create(
            empresa=self.empresa,
            periodo="Enero 2026",
            tipo="bimestral",
            rfc_empresa="CON111111AAA",
            nombre_razon_social="Contratista S.A."
        )
        
        # Trabajador 1 (Juan Perez)
        TrabajadorSUA.objects.create(
            importacion=self.importacion,
            nss="11111111111",
            rfc_curp="CURP11111111111111",
            nombre="Juan Perez",
            sdi=100.00
        )
        
        # Trabajador 2 (Pedro Gomez)
        TrabajadorSUA.objects.create(
            importacion=self.importacion,
            nss="22222222222",
            rfc_curp="CURP22222222222222",
            nombre="Pedro Gomez",
            sdi=200.00
        )
        
    def test_exportar_carga_trabajadores_resolves_beneficiary_from_contract(self):
        self.client.force_login(self.user)
        
        # Simular sesión de la empresa
        session = self.client.session
        session['empresa_id'] = self.empresa.id
        session.save()
        
        # Petición al reporte para el primer cuatrimestre de 2026
        url = reverse('exportar_carga_trabajadores', args=[self.contratista.id])
        response = self.client.get(url, {'cuatrimestre': 1, 'anio': 2026, 'formato': 'csv'})
        self.assertEqual(response.status_code, 200)
        
        content = response.content.decode('utf-8')
        # Ambos trabajadores (NSS 11111111111 y NSS 22222222222) deben estar presentes en el archivo CSV
        self.assertIn("11111111111", content)
        self.assertIn("22222222222", content)


class SisubTrabajadoresSortTest(TestCase):
    def setUp(self):
        self.empresa = Empresa.objects.create(
            nombre="Empresa de Prueba",
            subdominio="prueba",
            usuario_admin="admin",
            correo_contacto="prueba@test.com"
        )
        self.empresa.modulo_recursos_humanos = True
        self.empresa.save()

        self.user = User.objects.create_superuser(
            username="admin@prueba",
            email="admin@prueba.com",
            password="password"
        )
        
        self.contratista = Contratista.objects.create(
            empresa=self.empresa,
            nombre_razon_social="Contratista S.A.",
            rfc="CON111111AAA"
        )
        
        self.beneficiario = Beneficiario.objects.create(
            empresa=self.empresa,
            nombre_razon_social="Beneficiario S.A.",
            rfc="BEN111111AAA"
        )
        
        import datetime
        # Contrato con folio 10
        self.contrato_10 = Contrato.objects.create(
            empresa=self.empresa,
            contratista=self.contratista,
            beneficiario=self.beneficiario,
            folio="10",
            fecha_inicio=datetime.date(2026, 1, 1),
            fecha_fin=datetime.date(2026, 4, 30),
            monto_contrato=10000.00
        )
        
        # Contrato con folio 2
        self.contrato_2 = Contrato.objects.create(
            empresa=self.empresa,
            contratista=self.contratista,
            beneficiario=self.beneficiario,
            folio="2",
            fecha_inicio=datetime.date(2026, 1, 1),
            fecha_fin=datetime.date(2026, 4, 30),
            monto_contrato=10000.00
        )
        
        # Contrato con folio 1
        self.contrato_1 = Contrato.objects.create(
            empresa=self.empresa,
            contratista=self.contratista,
            beneficiario=self.beneficiario,
            folio="1",
            fecha_inicio=datetime.date(2026, 1, 1),
            fecha_fin=datetime.date(2026, 4, 30),
            monto_contrato=10000.00
        )
        
        self.emp1 = Empleado.objects.create(
            empresa=self.empresa,
            nombre="Juan",
            apellido_paterno="Perez",
            nss="11111111111",
            curp="CURP11111111111111",
            contratista=self.contratista
        )
        self.contrato_1.empleados.add(self.emp1)
        
        self.emp2 = Empleado.objects.create(
            empresa=self.empresa,
            nombre="Pedro",
            apellido_paterno="Gomez",
            nss="22222222222",
            curp="CURP22222222222222",
            contratista=self.contratista
        )
        self.contrato_2.empleados.add(self.emp2)

        self.emp10 = Empleado.objects.create(
            empresa=self.empresa,
            nombre="Maria",
            apellido_paterno="Lopez",
            nss="33333333333",
            curp="CURP33333333333333",
            contratista=self.contratista
        )
        self.contrato_10.empleados.add(self.emp10)

        # Crear Nominas para cada empleado
        Nomina.objects.create(
            empresa=self.empresa,
            nss="11111111111",
            nombre="Juan Perez",
            fecha_pago=datetime.date(2026, 1, 15),
            sueldo_gravado=Decimal("1000.00")
        )
        Nomina.objects.create(
            empresa=self.empresa,
            nss="22222222222",
            nombre="Pedro Gomez",
            fecha_pago=datetime.date(2026, 1, 15),
            sueldo_gravado=Decimal("2000.00")
        )
        Nomina.objects.create(
            empresa=self.empresa,
            nss="33333333333",
            nombre="Maria Lopez",
            fecha_pago=datetime.date(2026, 1, 15),
            sueldo_gravado=Decimal("3000.00")
        )

    def test_exportar_sisub_trabajadores_sorted_by_contract_number(self):
        self.client.force_login(self.user)
        session = self.client.session
        session['empresa_id'] = self.empresa.id
        session.save()
        
        url = reverse('exportar_sisub_trabajadores', args=[self.contratista.id])
        response = self.client.get(url, {'cuatrimestre': 1, 'anio': 2026, 'formato': 'csv'})
        self.assertEqual(response.status_code, 200)
        
        content = response.content.decode('utf-8')
        lines = content.split('\r\n')
        # Filter out empty lines
        lines = [line for line in lines if line.strip()]
        
        # Header is at index 0
        # Line 1 should be contract '1'
        # Line 2 should be contract '2'
        # Line 3 should be contract '10'
        
        # Let's check all data lines
        found_folios = []
        for line in lines[1:]:
            # contract folio is the 5th value (index 4) in comma separated values
            parts = line.split(',')
            if len(parts) > 4:
                found_folios.append(parts[4])
        
        self.assertEqual(found_folios, ["1", "2", "10"])


class CargarXMLDirectoTest(TestCase):
    def setUp(self):
        self.empresa = Empresa.objects.create(
            nombre="Empresa Prueba Upload",
            subdominio="pruebaupload",
            usuario_admin="admin",
            correo_contacto="upload@test.com"
        )
        self.empresa.modulo_recursos_humanos = True
        self.empresa.save()

        self.user = User.objects.create_superuser(
            username="admin@pruebaupload",
            email="admin@pruebaupload.com",
            password="password"
        )
        self.client.force_login(self.user)
        session = self.client.session
        session['empresa_id'] = self.empresa.id
        session.save()

        self.xml_content = """<cfdi:Comprobante xmlns:cfdi="http://www.sat.gob.mx/cfd/3" xmlns:tfd="http://www.sat.gob.mx/TimbreFiscalDigital" xmlns:nomina12="http://www.sat.gob.mx/nomina12" Version="3.3" Total="1500.00" Folio="888" Serie="N" Fecha="2026-07-15T12:30:00">
  <cfdi:Emisor Rfc="CON010101AAA"/>
  <cfdi:Receptor Rfc="EMP001122XX3" Nombre="JUAN VALENZUELA"/>
  <cfdi:Complemento>
    <tfd:TimbreFiscalDigital UUID="f08b8b2e-2222-3333-4444-555566667777" FechaTimbrado="2026-07-15T12:35:45"/>
    <nomina12:Nomina FechaPago="2026-07-15" FechaInicialPago="2026-07-01" FechaFinalPago="2026-07-15" NumDiasPagados="15.00" TipoNomina="O">
      <nomina12:Receptor Curp="JUAN000000HDFRXX02" NumSeguridadSocial="98765432101"/>
    </nomina12:Nomina>
  </cfdi:Complemento>
</cfdi:Comprobante>"""

    @patch('recursos_humanos.sat_service.CFDI')
    def test_cargar_xml_directo_xml_unico(self, mock_cfdi):
        mock_cfdi.from_string.return_value = "mock_cfdi"
        from django.core.files.uploadedfile import SimpleUploadedFile
        import datetime
        
        xml_file = SimpleUploadedFile("nomina_directo.xml", self.xml_content.encode('utf-8'), content_type="text/xml")
        url = reverse('cargar_xml_directo_ajax')
        
        response = self.client.post(url, {
            'archivos': [xml_file],
            'estatus': 'vigente',
            'sucursal': ''
        })
        
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data['status'], 'success')
        self.assertIn("Se cargaron e integraron con éxito 1", data['message'])
        
        # Verify it exists in database
        nomina = Nomina.objects.get(uuid="f08b8b2e-2222-3333-4444-555566667777", empresa=self.empresa)
        self.assertEqual(nomina.folio, "888")
        self.assertEqual(nomina.estado, "vigente")
        self.assertIsNotNone(nomina.fecha_emision)
        self.assertIsNotNone(nomina.fecha_certificacion)
        self.assertEqual(nomina.fecha_pago, datetime.date(2026, 7, 15))
        self.assertEqual(nomina.fecha_inicial_pago, datetime.date(2026, 7, 1))
        self.assertEqual(nomina.fecha_final_pago, datetime.date(2026, 7, 15))

    @patch('recursos_humanos.sat_service.CFDI')
    def test_cargar_xml_directo_zip(self, mock_cfdi):
        mock_cfdi.from_string.return_value = "mock_cfdi"
        from django.core.files.uploadedfile import SimpleUploadedFile
        import zipfile
        import io
        import datetime
        
        zip_buffer = io.BytesIO()
        with zipfile.ZipFile(zip_buffer, 'w') as z:
            z.writestr("dentro_del_lote.xml", self.xml_content)
        zip_buffer.seek(0)
        
        zip_file = SimpleUploadedFile("lote_nominas.zip", zip_buffer.read(), content_type="application/zip")
        url = reverse('cargar_xml_directo_ajax')
        
        response = self.client.post(url, {
            'archivos': [zip_file],
            'estatus': 'cancelado',
            'sucursal': ''
        })
        
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data['status'], 'success')
        self.assertIn("Se cargaron e integraron con éxito 1", data['message'])
        
        # Verify it exists in database
        nomina = Nomina.objects.get(uuid="f08b8b2e-2222-3333-4444-555566667777", empresa=self.empresa)
        self.assertEqual(nomina.estado, "cancelado")
        self.assertIsNotNone(nomina.fecha_emision)
        self.assertIsNotNone(nomina.fecha_certificacion)
        self.assertEqual(nomina.fecha_pago, datetime.date(2026, 7, 15))
        self.assertEqual(nomina.fecha_inicial_pago, datetime.date(2026, 7, 1))
        self.assertEqual(nomina.fecha_final_pago, datetime.date(2026, 7, 15))


class SUAFilterDropdownsTest(TestCase):
    def setUp(self):
        self.empresa = Empresa.objects.create(
            nombre="Empresa SUA Test",
            subdominio="suatest",
            usuario_admin="admin",
            correo_contacto="sua@test.com"
        )
        self.empresa.modulo_recursos_humanos = True
        self.empresa.save()

        self.user = User.objects.create_superuser(
            username="admin@suatest",
            email="admin@suatest.com",
            password="password"
        )
        self.client.force_login(self.user)
        session = self.client.session
        session['empresa_id'] = self.empresa.id
        session.save()

        from recursos_humanos.models import ImportacionSUA
        ImportacionSUA.objects.create(
            empresa=self.empresa,
            registro_patronal="RP999888777",
            nombre_razon_social="RAZON SOCIAL DE PRUEBA",
            periodo="Julio-2026",
            tipo="mensual"
        )

    def test_lista_sua_dropdown_context(self):
        url = reverse('lista_sua')
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        
        # Verify that unique registry patronals and business names are in the context
        self.assertIn('reg_patronales_unicos', response.context)
        self.assertIn('razones_sociales_unicas', response.context)
        
        reg_pats = list(response.context['reg_patronales_unicos'])
        razones = list(response.context['razones_sociales_unicas'])
        
        self.assertEqual(reg_pats, ["RP999888777"])
        self.assertEqual(razones, ["RAZON SOCIAL DE PRUEBA"])


class BeneficiarioFilterDropdownsTest(TestCase):
    def setUp(self):
        self.empresa = Empresa.objects.create(
            nombre="Empresa Beneficiario Test",
            subdominio="bentest",
            usuario_admin="admin",
            correo_contacto="ben@test.com"
        )
        self.empresa.modulo_recursos_humanos = True
        self.empresa.save()

        self.user = User.objects.create_superuser(
            username="admin@bentest",
            email="admin@bentest.com",
            password="password"
        )
        self.client.force_login(self.user)
        session = self.client.session
        session['empresa_id'] = self.empresa.id
        session.save()

        from recursos_humanos.models import Beneficiario
        Beneficiario.objects.create(
            empresa=self.empresa,
            nombre_razon_social="RAZON SOCIAL BENEFICIARIO",
            rfc="RFCBEN9998887",
            registro_patronal="RPBEN888999"
        )

    def test_lista_beneficiarios_dropdown_context(self):
        url = reverse('lista_beneficiarios')
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        
        # Verify context options
        self.assertIn('razones_sociales_unicas', response.context)
        self.assertIn('rfcs_unicos', response.context)
        self.assertIn('reg_patronales_unicos', response.context)
        
        razones = list(response.context['razones_sociales_unicas'])
        rfcs = list(response.context['rfcs_unicos'])
        reg_pats = list(response.context['reg_patronales_unicos'])
        
        self.assertEqual(razones, ["RAZON SOCIAL BENEFICIARIO"])
        self.assertEqual(rfcs, ["RFCBEN9998887"])
        self.assertEqual(reg_pats, ["RPBEN888999"])


class ContratistaFilterDropdownsTest(TestCase):
    def setUp(self):
        self.empresa = Empresa.objects.create(
            nombre="Empresa Contratista Test",
            subdominio="conttest",
            usuario_admin="admin",
            correo_contacto="con@test.com"
        )
        self.empresa.modulo_recursos_humanos = True
        self.empresa.save()

        self.user = User.objects.create_superuser(
            username="admin@conttest",
            email="admin@conttest.com",
            password="password"
        )
        self.client.force_login(self.user)
        session = self.client.session
        session['empresa_id'] = self.empresa.id
        session.save()

        from recursos_humanos.models import Contratista
        Contratista.objects.create(
            empresa=self.empresa,
            nombre_razon_social="RAZON SOCIAL CONTRATISTA",
            rfc="RFCCON9998887",
            registro_patronal="RPCON888999",
            correo="con@test.com"
        )

    def test_lista_contratistas_dropdown_context(self):
        url = reverse('lista_contratistas')
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        
        # Verify context options
        self.assertIn('razones_sociales_unicas', response.context)
        self.assertIn('rfcs_unicos', response.context)
        self.assertIn('reg_patronales_unicos', response.context)
        
        razones = list(response.context['razones_sociales_unicas'])
        rfcs = list(response.context['rfcs_unicos'])
        reg_pats = list(response.context['reg_patronales_unicos'])
        
        self.assertEqual(razones, ["RAZON SOCIAL CONTRATISTA"])
        self.assertEqual(rfcs, ["RFCCON9998887"])
        self.assertEqual(reg_pats, ["RPCON888999"])


class DefaultSucursalFilterTest(TestCase):
    def setUp(self):
        from preferencias.models import Sucursal
        self.empresa = Empresa.objects.create(
            nombre="Empresa Sucursal Test",
            subdominio="suctest",
            usuario_admin="admin",
            correo_contacto="suc@test.com"
        )
        self.empresa.modulo_recursos_humanos = True
        self.empresa.save()

        self.sucursal = Sucursal.objects.create(
            empresa=self.empresa,
            nombre="Sucursal Norte"
        )

        self.user = User.objects.create_superuser(
            username="admin@suctest",
            email="admin@suctest.com",
            password="password"
        )
        self.client.force_login(self.user)
        session = self.client.session
        session['empresa_id'] = self.empresa.id
        session['sucursal_id'] = self.sucursal.id
        session.save()

    def test_default_sucursal_in_all_submodules(self):
        views_to_test = [
            'lista_empleados',
            'lista_contratos',
            'lista_contratistas',
            'lista_beneficiarios',
            'lista_sua',
            'lista_nomina'
        ]

        for view_name in views_to_test:
            url = reverse(view_name)
            
            # Scenario A: First load (no sucursal parameter in query string)
            response = self.client.get(url)
            self.assertEqual(response.status_code, 200, f"Error in view {view_name}")
            self.assertEqual(
                response.context['filtros']['sucursal'],
                str(self.sucursal.id),
                f"Sucursal was not defaulted to session sucursal in view {view_name}"
            )

            # Scenario B: Explicit selection of "Todas" (sucursal="" in query string)
            response = self.client.get(url, {'sucursal': ''})
            self.assertEqual(response.status_code, 200, f"Error in view {view_name} with explicit Todas")
            self.assertEqual(
                response.context['filtros']['sucursal'],
                '',
                f"Explicit 'Todas' was ignored in view {view_name}"
            )


class ProveedoresListaContratistasTest(TestCase):
    def setUp(self):
        self.empresa = Empresa.objects.create(
            nombre="Empresa Proveedores Test",
            subdominio="provtest",
            usuario_admin="admin",
            correo_contacto="prov@test.com"
        )
        self.empresa.modulo_recursos_humanos = True
        self.empresa.save()

        self.user = User.objects.create_superuser(
            username="admin@provtest",
            email="admin@provtest.com",
            password="password"
        )
        self.client.force_login(self.user)
        session = self.client.session
        session['empresa_id'] = self.empresa.id
        session.save()

    def test_proveedor_model_save_autogenerates_domicilio(self):
        from recursos_humanos.models import ProveedorRH
        p = ProveedorRH.objects.create(
            empresa=self.empresa,
            clave="PROV-001",
            nombre_razon_social="PROVEEDOR DE PRUEBA S.A.",
            rfc="PRV880808XYZ",
            calle="Av. Juárez",
            num_ext="123",
            num_int="4B",
            colonia="Centro",
            cp="06000",
            municipio_alcaldia="Cuauhtémoc"
        )
        # Verify autogenerated domicilio
        self.assertEqual(p.domicilio, "Av. Juárez, No. 123, Int. 4B, Col. Centro, Cuauhtémoc, C.P. 06000")

    def test_lista_contratistas_with_tipo_proveedor(self):
        from recursos_humanos.models import ProveedorRH, Contratista
        # Create a contractor
        contratista = Contratista.objects.create(
            empresa=self.empresa,
            rfc="CON121212ABC",
            nombre_razon_social="CONTRATISTA PRUEBA"
        )
        # Create a provider
        ProveedorRH.objects.create(
            empresa=self.empresa,
            contratista=contratista,
            clave="PROV-002",
            nombre_razon_social="DISTRIBUIDORA DE INSUMOS",
            rfc="DIS990909ABC",
            calle="Reforma",
            cp="06500"
        )

        url = reverse('lista_contratistas')
        response = self.client.get(url, {'tipo': 'proveedor', 'contratista_id': contratista.id})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['tipo_vista'], 'proveedor')
        self.assertIn('razones_sociales_unicas', response.context)
        self.assertIn('rfcs_unicos', response.context)
        
        # Verify provider in list
        razones = list(response.context['razones_sociales_unicas'])
        self.assertEqual(razones, ["DISTRIBUIDORA DE INSUMOS"])


from django.core.files.uploadedfile import SimpleUploadedFile
from django.core import mail
from recursos_humanos.models import DocumentacionProveedor, ProveedorRH

class ProveedorDocumentoNotificationTest(TestCase):
    def setUp(self):
        self.empresa = Empresa.objects.create(
            nombre="Empresa Test",
            subdominio="test",
            usuario_admin="admin",
            correo_contacto="test@test.com"
        )
        self.contratista = Contratista.objects.create(
            empresa=self.empresa,
            rfc="CON121212ABC",
            nombre_razon_social="Contratista S.A.",
            correo="contratista@test.com"
        )
        self.proveedor = ProveedorRH.objects.create(
            empresa=self.empresa,
            contratista=self.contratista,
            nombre_razon_social="Proveedor Express",
            rfc="PROV990909XYZ",
            correo="proveedor@test.com"
        )
        # Crear usuario asociado al proveedor
        self.prov_user = User.objects.create_user(
            username="prov_express@test",
            password="password",
            email="proveedor@test.com"
        )
        self.proveedor.usuario = self.prov_user
        self.proveedor.save()

    def test_subir_documento_envia_notificacion_a_contratista(self):
        self.client.login(username="prov_express@test", password="password")
        
        # Simular sesión de sucursal
        session = self.client.session
        session['sucursal_id'] = None
        session.save()

        test_file = SimpleUploadedFile(
            "repse.pdf",
            b"pdf-content-bytes",
            content_type="application/pdf"
        )
        
        post_data = {
            'nombre_documento': 'REPSE_VIGENTE',
            'mes': '8',
            'anio': '2026',
            'archivo': test_file
        }

        # Subir documento
        response = self.client.post(
            reverse('subir_documento_proveedor_ajax', args=[self.proveedor.id]),
            post_data
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()['success'])

        # Verificar que el documento se guardó
        self.assertTrue(DocumentacionProveedor.objects.filter(
            proveedor=self.proveedor,
            nombre_documento='REPSE_VIGENTE',
            mes=8,
            anio=2026
        ).exists())

        # Verificar que se envió la notificación por correo
        self.assertEqual(len(mail.outbox), 1)
        sent_email = mail.outbox[0]
        self.assertEqual(sent_email.to, ["contratista@test.com"])
        self.assertIn("Contratista S.A.", sent_email.body)
        self.assertIn("Proveedor Express", sent_email.body)
        self.assertIn("Constancia/registro REPSE vigente", sent_email.body)
        self.assertIn("8/2026", sent_email.body)
        self.assertEqual(len(sent_email.attachments), 1)
        self.assertTrue("repse" in sent_email.attachments[0][0])


class ContratistaSMTPTest(TestCase):
    def setUp(self):
        self.empresa = Empresa.objects.create(
            nombre="Empresa Test SMTP",
            subdominio="test_smtp",
            usuario_admin="admin",
            correo_contacto="test@test.com"
        )
        self.user = User.objects.create_superuser(
            username="admin@test_smtp",
            password="password",
            email="admin@test.com"
        )
        self.client.login(username="admin@test_smtp", password="password")
        
        # Simular sesión de sucursal
        session = self.client.session
        session['sucursal_id'] = None
        session.save()

    def test_crear_y_obtener_contratista_con_smtp(self):
        post_data = {
            'clave': 'CONT-SMTP-01',
            'rfc': 'CON121212XYZ',
            'nombre_razon_social': 'Contratista SMTP S.A.',
            'correo': 'correo_contratista@test.com',
            'smtp_host': 'smtp.testmail.com',
            'smtp_port': '587',
            'smtp_user': 'sender@testmail.com',
            'smtp_password': 'smtp_secret_pass',
            'use_tls': 'true',
            'use_ssl': 'false',
            'nombre_remitente': 'Notificaciones Test',
            'email_remitente': 'noreply@testmail.com',
            'email_notificacion_1': 'notif_principal@testmail.com',
            'email_notificacion_2': 'copia1@testmail.com',
            'email_notificacion_3': 'copia2@testmail.com',
        }
        
        # Crear Contratista via AJAX
        response = self.client.post(reverse('crear_contratista_ajax'), post_data)
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()['success'])
        
        # Obtener el contratista creado
        from recursos_humanos.models import Contratista, ContratistaCorreoSMTP
        contratista = Contratista.objects.get(rfc='CON121212XYZ')
        self.assertEqual(contratista.nombre_razon_social, 'Contratista SMTP S.A.')
        
        # Verificar que la configuración SMTP y destinatarios se hayan guardado
        smtp_config = ContratistaCorreoSMTP.objects.get(contratista=contratista)
        self.assertEqual(smtp_config.smtp_host, 'smtp.testmail.com')
        self.assertEqual(smtp_config.smtp_user, 'sender@testmail.com')
        self.assertEqual(smtp_config.smtp_password, 'smtp_secret_pass')
        self.assertTrue(smtp_config.use_tls)
        self.assertFalse(smtp_config.use_ssl)
        self.assertEqual(smtp_config.email_notificacion_1, 'notif_principal@testmail.com')
        self.assertEqual(smtp_config.email_notificacion_2, 'copia1@testmail.com')
        self.assertEqual(smtp_config.email_notificacion_3, 'copia2@testmail.com')
        
        # Probar el JSON de obtención
        response_json = self.client.get(reverse('obtener_contratista_json', args=[contratista.id]))
        self.assertEqual(response_json.status_code, 200)
        data = response_json.json()['data']
        self.assertEqual(data['smtp_host'], 'smtp.testmail.com')
        self.assertEqual(data['smtp_user'], 'sender@testmail.com')
        self.assertEqual(data['nombre_remitente'], 'Notificaciones Test')
        self.assertEqual(data['email_notificacion_1'], 'notif_principal@testmail.com')
        self.assertEqual(data['email_notificacion_2'], 'copia1@testmail.com')
        self.assertEqual(data['email_notificacion_3'], 'copia2@testmail.com')

    def test_subir_documento_envia_a_destinatarios_y_cc(self):
        from recursos_humanos.models import Contratista, ContratistaCorreoSMTP, ProveedorRH, DocumentacionProveedor
        from django.core.files.uploadedfile import SimpleUploadedFile
        from django.core import mail

        mail.outbox = []

        contratista = Contratista.objects.create(
            empresa=self.empresa,
            rfc="CON999999ABC",
            nombre_razon_social="Contratista Destinatarios S.A.",
            correo="general_contratista@test.com"
        )
        ContratistaCorreoSMTP.objects.create(
            contratista=contratista,
            email_notificacion_1="principal@notificaciones.com",
            email_notificacion_2="cc1@notificaciones.com",
            email_notificacion_3="cc2@notificaciones.com",
        )

        proveedor = ProveedorRH.objects.create(
            empresa=self.empresa,
            nombre_razon_social="Proveedor Con Notificaciones",
            rfc="PROV999999XYZ",
            contratista=contratista,
            creado_por=self.user
        )

        test_file = SimpleUploadedFile("doc_prueba.pdf", b"dummy content", content_type="application/pdf")
        post_data = {
            'nombre_documento': 'REPSE_VIGENTE',
            'mes': '8',
            'anio': '2026',
            'archivo': test_file
        }

        response = self.client.post(
            reverse('subir_documento_proveedor_ajax', args=[proveedor.id]),
            post_data
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()['success'])

        self.assertEqual(len(mail.outbox), 1)
        sent = mail.outbox[0]
        self.assertEqual(sent.to, ["principal@notificaciones.com"])
        self.assertEqual(sent.cc, ["cc1@notificaciones.com", "cc2@notificaciones.com"])
        self.assertIn("Contratista Destinatarios S.A.", sent.body)
        self.assertIn("Proveedor Con Notificaciones", sent.body)








