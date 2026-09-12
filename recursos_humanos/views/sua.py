from django.shortcuts import render, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse, HttpResponse
from django.core.paginator import Paginator
from django.views.decorators.http import require_POST
from django.db import transaction
from django.db.models import Q
from decimal import Decimal
import re
import csv
import openpyxl
from datetime import datetime

try:
    import pdfplumber
except ImportError:
    pdfplumber = None

from ..models import Empleado, Contrato, Contratista, Beneficiario, ImportacionSUA, TrabajadorSUA
from preferencias.models import Sucursal
from preferencias.permissions import require_hr_permission, user_has_hr_permission
from .utils import get_empresa_actual, limpiar_basura_header
from notificaciones.utils import crear_notificacion

@login_required(login_url='/login/')
@require_hr_permission('sua', 'ver')
def lista_sua(request):
    empresa_actual = get_empresa_actual(request)
    importaciones = ImportacionSUA.objects.filter(empresa=empresa_actual).select_related('sucursal').order_by('-fecha_importacion')
    
    q_reg_pat = request.GET.get('reg_patronal', '')
    q_razon = request.GET.get('razon_social', '')
    q_periodo = request.GET.get('periodo', '')
    sucursal_id = request.GET.get('sucursal')
    if sucursal_id is None:
        sucursal_id = str(request.session.get('sucursal_id') or '')

    if q_reg_pat:
        importaciones = importaciones.filter(registro_patronal__icontains=q_reg_pat)
    if q_razon:
        importaciones = importaciones.filter(nombre_razon_social__icontains=q_razon)
    if q_periodo:
        importaciones = importaciones.filter(periodo__icontains=q_periodo)
    if sucursal_id:
        importaciones = importaciones.filter(sucursal_id=sucursal_id)

    sucursales = Sucursal.objects.filter(empresa=empresa_actual).order_by('nombre')
    
    # Obtener valores únicos para los filtros desplegables
    reg_patronales_unicos = ImportacionSUA.objects.filter(empresa=empresa_actual).exclude(registro_patronal='').values_list('registro_patronal', flat=True).distinct().order_by('registro_patronal')
    razones_sociales_unicas = ImportacionSUA.objects.filter(empresa=empresa_actual).exclude(nombre_razon_social='').values_list('nombre_razon_social', flat=True).distinct().order_by('nombre_razon_social')
    
    # Ordenar importaciones cronológicamente (de la más reciente a la más vieja)
    importaciones_list = list(importaciones)
    meses_es = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]
    
    def get_periodo_sort_key(imp):
        parts = imp.periodo.split('-')
        if len(parts) == 2:
            mes_str, anio_str = parts
            try:
                anio = int(anio_str)
                mes = meses_es.index(mes_str.capitalize()) + 1
                return (anio, mes)
            except (ValueError, IndexError):
                pass
        return (imp.fecha_importacion.year, imp.fecha_importacion.month)
        
    importaciones_list.sort(key=get_periodo_sort_key, reverse=True)

    # PAGINACIÓN
    paginator = Paginator(importaciones_list, 20)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    return render(request, 'recursos_humanos/lista_sua.html', {
        'page_obj': page_obj, 
        'sucursales': sucursales, 
        'empresa': empresa_actual,
        'reg_patronales_unicos': reg_patronales_unicos,
        'razones_sociales_unicas': razones_sociales_unicas,
        'filtros': {
            'reg_patronal': q_reg_pat, 
            'razon_social': q_razon, 
            'periodo': q_periodo, 
            'sucursal': sucursal_id
        }
    })

@login_required(login_url='/login/')
@require_POST
@require_hr_permission('sua', 'importar', json_response=True)
def importar_sua_ajax(request):
    empresa_actual = get_empresa_actual(request)
    pdf_file = request.FILES.get('archivo_sua')
    tipo_importacion = request.POST.get('tipo', 'bimestral')
    if not pdf_file: return JsonResponse({'success': False, 'error': 'No se proporcionó ningún archivo.'})
    try:
        if not pdfplumber: return JsonResponse({'success': False, 'error': 'Librería pdfplumber no instalada.'})
        
        text_lines = []
        with pdfplumber.open(pdf_file) as pdf:
            for page in pdf.pages:
                text = page.extract_text(layout=True)
                if text:
                    text_lines.extend(text.split('\n'))
        
        full_text = "\n".join(text_lines)
        
        reg_pat_val, rfc_emp_val, area_val = "", "", ""
        nom_razon_val, deleg_val = "", ""
        actividad_val, subdeleg_val = "", ""
        domicilio_val, mun_alc_val = "", ""
        cp_val, entidad_val, prima_val = "", "", ""
        periodo_val = "Desconocido"

        for line in text_lines:
            if "Registro Patronal:" in line:
                m_rp = re.search(r'Registro Patronal:\s*([\w-]+)', line, re.I)
                m_rfc = re.search(r'RFC:\s*([\w\d-]+)', line, re.I)
                m_area = re.search(r'Area Geográfica:\s*(.*?)(?=\s{2,}|Delegación|$)', line, re.I)
                if m_rp: reg_pat_val = m_rp.group(1).strip()
                if m_rfc: rfc_emp_val = m_rfc.group(1).strip()
                if m_area: area_val = m_area.group(1).strip()
            
            if "Nombre o Razón Social:" in line:
                m_nom = re.search(r'Nombre o Razón Social:\s*(.*?)(?=\s{2,}|Delegación|Convenio|Aportación|V \d|$)', line, re.I)
                m_del = re.search(r'Delegación IMSS:\s*(.*?)(?=\s{2,}|Fecha|Convenio|Aportación|V \d|$)', line, re.I)
                if m_nom: nom_razon_val = m_nom.group(1).strip()
                if m_del: deleg_val = m_del.group(1).strip()

            if "Actividad:" in line:
                m_act = re.search(r'Actividad:\s*(.*?)(?=\s{2,}|Subdelegación|Convenio|Aportación|V \d|$)', line, re.I)
                m_sub = re.search(r'SubDelegación IMSS:\s*(.*?)(?=\s{2,}|Area|Convenio|Aportación|V \d|$)', line, re.I)
                if m_act: actividad_val = m_act.group(1).strip()
                if m_sub: subdeleg_val = m_sub.group(1).strip()

            if "Domicilio:" in line:
                m_dom = re.search(r'Domicilio:\s*(.*?)(?=\s{2,}|Pob|Convenio|Aportación|V \d|$)', line, re.I)
                m_mun = re.search(r'Pob\., Mun\. / Alcaldía:\s*(.*?)(?=\s{2,}|Prima|Convenio|Aportación|V \d|$)', line, re.I)
                if m_dom: domicilio_val = m_dom.group(1).strip()
                if m_mun: mun_alc_val = m_mun.group(1).strip()

            if "Código Postal:" in line:
                m_cp = re.search(r'Código Postal:\s*(\d+)', line, re.I)
                m_ent = re.search(r'Entidad:\s*(.*?)(?=\s{2,}|Prima|Convenio|Aportación|V \d|$)', line, re.I)
                m_pri = re.search(r'Prima de R\.T\.\s*([\d\.,%]+)', line, re.I)
                if m_cp: cp_val = m_cp.group(1).strip()
                if m_ent: entidad_val = m_ent.group(1).strip()
                if m_pri: prima_val = m_pri.group(1).strip()
            
            if "Proceso:" in line:
                m_per = re.search(r'(?:Período|Bimestre)\s+de\s+Proceso:\s*([\w\d-]+)', line, re.I)
                if m_per: periodo_val = m_per.group(1).strip()

        periodo_final = limpiar_basura_header(periodo_val).strip()
        razon_final = nom_razon_val.strip()
        reg_pat_final = reg_pat_val.strip()
        
        if ImportacionSUA.objects.filter(
            empresa=empresa_actual, periodo=periodo_final,
            nombre_razon_social=razon_final, registro_patronal=reg_pat_final,
            tipo=tipo_importacion
        ).exists():
            return JsonResponse({'success': False, 'error': f'Error: Esta cédula ya fue registrada (Periodo: {periodo_final}, Empresa: {razon_final}, Tipo: {tipo_importacion.upper()}).'})

        total_reporte = 0
        m_total_rep = re.search(r'Total\s+de\s+cotizaciones:\s*(\d+)', full_text, re.I)
        if m_total_rep: total_reporte = int(m_total_rep.group(1))

        sucursal_id = request.session.get('sucursal_id')
        created_count = 0; nss_encontrados = set()

        with transaction.atomic():
            importacion = ImportacionSUA.objects.create(
                empresa=empresa_actual, sucursal_id=sucursal_id,
                registro_patronal=reg_pat_val, rfc_empresa=rfc_emp_val,
                nombre_razon_social=nom_razon_val, actividad=limpiar_basura_header(actividad_val),
                domicilio=limpiar_basura_header(domicilio_val), cp=limpiar_basura_header(cp_val),
                entidad=limpiar_basura_header(entidad_val), area_geografica=limpiar_basura_header(area_val),
                delegacion_imss=limpiar_basura_header(deleg_val), subdelegacion_imss=limpiar_basura_header(subdeleg_val),
                municipio_alcaldia=limpiar_basura_header(mun_alc_val), prima_rt=prima_val,
                periodo=periodo_final, tipo=tipo_importacion,
                creado_por=request.user
            )

            current_worker_info = None; stop_workers = False
            for line in text_lines:
                l_clean = line.strip()
                if not l_clean: continue
                if "TOTAL DE COTIZACIONES" in l_clean.upper() or "TOTALES" in l_clean.upper() or re.search(r'([_-]\s?){15,}', l_clean):
                    current_worker_info = None; stop_workers = True; continue
                if stop_workers: continue

                nss_match = re.search(r'(\d{2}-?\d{2}-?\d{2}-?\d{4}-?\d)', l_clean)
                if nss_match:
                    nss_val = nss_match.group(1); current_worker_info = None 
                    remainder = l_clean.split(nss_val)[-1].strip()
                    m_rfc = re.search(r'([A-Z]{3,4}[0-9]{6}[A-Z0-9]{0,9})', remainder)
                    
                    if m_rfc:
                        rfc = m_rfc.group(1); nombre = remainder[:m_rfc.start()].strip()
                        post_rfc = remainder[m_rfc.end():].strip()
                        m_data_start = re.search(r'(\d{1,2})\s+([\d\.,]+)', post_rfc)
                        if m_data_start:
                            ubic = post_rfc[:m_data_start.start()].strip() or "-"; l_clean = post_rfc[m_data_start.start():].strip()
                        else:
                            ubic = post_rfc or "-"; l_clean = ""
                        if nombre and rfc:
                            current_worker_info = {'nss': nss_val, 'nombre': nombre, 'rfc': rfc, 'clave_u': ubic}
                            nss_encontrados.add(nss_val)
                    else:
                        parts = [p.strip() for p in re.split(r'\s+', remainder) if p.strip()]
                        if len(parts) >= 2:
                            nombre = parts[0]; rfc = parts[1]; ubic = " ".join(parts[2:]) if len(parts) > 2 else "-"
                            current_worker_info = {'nss': nss_val, 'nombre': nombre, 'rfc': rfc, 'clave_u': ubic}
                            nss_encontrados.add(nss_val); l_clean = ""
                    if not l_clean: continue

                if current_worker_info:
                    m_header = re.match(r'^([^0-9\s,]{2,})?\s*(\d{2}/\d{2}/\d{4})?\s*(.*)', l_clean, re.I)
                    clave_mov = m_header.group(1).strip().rstrip(',') if m_header and m_header.group(1) else '-'
                    fecha_mov = m_header.group(2) or '' if m_header else ''; resto_linea = m_header.group(3).strip() if m_header else l_clean
                    raw_tokens = re.findall(r'\d{2}/\d{2}/\d{4}|(?:FD|\$|%)\s*[\d\.,]+%?|[A-Z]{3}|[\d\.,]+%?|FD|%|\$', resto_linea)
                    raw_tokens = [t.strip() for t in raw_tokens if t.strip()]
                    
                    def merge_tokens_sua(items):
                        res = []; skip = False
                        for i in range(len(items)):
                            if skip: skip = False; continue
                            token = items[i]
                            if token in ['FD', '$', '%'] and (i + 1) < len(items):
                                next_t = items[i+1]
                                if re.match(r'^[\d\.,]+', next_t): res.append(f"{token} {next_t}"); skip = True; continue
                            res.append(token)
                        return res

                    tokens = merge_tokens_sua(raw_tokens); tokens_clean = [t.replace(',', '') for t in tokens]
                    
                    def clean_dec(val):
                        if not val or val == '-': return Decimal('0')
                        c = re.sub(r'[^\d.]', '', str(val)); return Decimal(c) if c else Decimal('0')

                    try:
                        num_idx = 0
                        if tokens_clean[num_idx] == 'FD': num_idx += 1
                        dias_val = int(float(clean_dec(tokens_clean[num_idx])))
                        tiene_sdi = len(tokens_clean) > (num_idx + 1) and re.match(r'^[^\d]*[\d\.]+[^\d]*$', tokens_clean[num_idx+1]) and float(clean_dec(tokens_clean[num_idx+1])) > 0
                        es_movimiento_datos = (dias_val <= 99 and tiene_sdi)
                        if es_movimiento_datos:
                            tokens = tokens[num_idx:]; tokens_clean = tokens_clean[num_idx:]
                    except: es_movimiento_datos = False

                    if (clave_mov.lower() in ['baja', 'reingreso', 'modificación', 'alta'] and fecha_mov != "") or es_movimiento_datos:
                        trabajador_data = {'importacion': importacion, 'nss': current_worker_info['nss'], 'nombre': current_worker_info['nombre'], 'rfc_curp': current_worker_info['rfc'], 'clave_ubicacion': current_worker_info['clave_u'], 'clave_mov': clave_mov, 'fecha_mov': fecha_mov, 'dias': 0, 'sdi': 0, 'licencias': 0, 'incapacidades': 0, 'ausentismos': 0, 'total_general': 0}
                        if es_movimiento_datos:
                            try:
                                trabajador_data.update({'dias': int(float(clean_dec(tokens_clean[0]))), 'sdi': clean_dec(tokens_clean[1]), 'licencias': int(float(clean_dec(tokens_clean[2]))), 'incapacidades': int(float(clean_dec(tokens_clean[3]))), 'ausentismos': int(float(clean_dec(tokens_clean[4])))})
                                if tipo_importacion == 'mensual':
                                    if len(tokens_clean) >= 19:
                                        trabajador_data.update({'cuota_fija': clean_dec(tokens_clean[5]), 'excedente_patronal': clean_dec(tokens_clean[6]), 'excedente_obrera': clean_dec(tokens_clean[7]), 'prestaciones_dinero_patronal': clean_dec(tokens_clean[8]), 'prestaciones_dinero_obrera': clean_dec(tokens_clean[9]), 'gastos_medicos_patronal': clean_dec(tokens_clean[10]), 'gastos_medicos_obrera': clean_dec(tokens_clean[11]), 'riesgo_trabajo_cuota': clean_dec(tokens_clean[12]), 'invalidez_vida_patronal': clean_dec(tokens_clean[13]), 'invalidez_vida_obrera': clean_dec(tokens_clean[14]), 'guarderias_ps': clean_dec(tokens_clean[15]), 'imss_patronal': clean_dec(tokens_clean[16]), 'imss_obrera': clean_dec(tokens_clean[17]), 'imss_subtotal': clean_dec(tokens_clean[18]), 'total_general': clean_dec(tokens_clean[18])})
                                else:
                                    pivot_idx = -1
                                    for idx, t in enumerate(tokens):
                                        if any(pref in t for pref in ['FD', '$', '%']) and re.search(r'\d', t): pivot_idx = idx; break
                                    if pivot_idx != -1:
                                        try:
                                            trabajador_data.update({'retiro': clean_dec(tokens_clean[pivot_idx - 5]), 'patronal': clean_dec(tokens_clean[pivot_idx - 4]), 'obrera': clean_dec(tokens_clean[pivot_idx - 3]), 'subtotal': clean_dec(tokens_clean[pivot_idx - 2])})
                                            ap_pat = clean_dec(tokens_clean[pivot_idx - 1]); amort = clean_dec(tokens_clean[pivot_idx + 1])
                                            trabajador_data.update({'aportacion_patronal': ap_pat, 'tipo_valor_infonavit': tokens[pivot_idx], 'amortizacion': amort, 'suma_infonavit': ap_pat + amort})
                                            if len(tokens) > pivot_idx + 2:
                                                for extra_idx in range(pivot_idx + 3, len(tokens)):
                                                    t_extra = tokens[extra_idx]
                                                    if re.match(r'^\d{8,11}$', t_extra): trabajador_data['cred_vivienda'] = t_extra
                                                    elif re.match(r'^[A-Z]{3}$', t_extra): trabajador_data['tipo_mov_credito'] = t_extra
                                                    elif re.match(r'^\d{2}/\d{2}/\d{4}$', t_extra): trabajador_data['fecha_mov_credito'] = t_extra
                                        except: pass
                                    else:
                                        try:
                                            if len(tokens_clean) >= 9: trabajador_data.update({'retiro': clean_dec(tokens_clean[5]), 'patronal': clean_dec(tokens_clean[6]), 'obrera': clean_dec(tokens_clean[7]), 'subtotal': clean_dec(tokens_clean[8])})
                                            if len(tokens_clean) >= 10: ap_pat = clean_dec(tokens_clean[9]); trabajador_data.update({'aportacion_patronal': ap_pat, 'suma_infonavit': ap_pat, 'amortizacion': Decimal('0'), 'tipo_valor_infonavit': '-'})
                                        except: pass
                                    trabajador_data['total_general'] = clean_dec(trabajador_data.get('subtotal', 0)) + clean_dec(trabajador_data.get('suma_infonavit', 0))
                            except: pass
                        TrabajadorSUA.objects.create(**trabajador_data); created_count += 1
                    else:
                        if not nss_match: current_worker_info = None

            if created_count == 0: raise Exception("No se detectaron trabajadores válidos.")
            unique_count = len(nss_encontrados)
            msg_val = f" Advertencia: Se detectaron {unique_count} trabajadores únicos pero el reporte indica un total de {total_reporte}." if total_reporte > 0 and unique_count != total_reporte else ""

        crear_notificacion(
            empresa=empresa_actual,
            actor=request.user,
            mensaje=f'importó reporte SUA {importacion.periodo}',
            link='/recursos-humanos/sua/',
            propietario=request.user
        )
        return JsonResponse({'success': True, 'message': f'Importación exitosa: {created_count} registros procesados.{msg_val}'})
    except Exception as e: return JsonResponse({'success': False, 'error': str(e)})

@login_required(login_url='/login/')
@require_hr_permission('sua', 'ver', json_response=True)
def obtener_registro_sua_json(request, id):
    empresa_actual = get_empresa_actual(request)
    try:
        imp = ImportacionSUA.objects.get(id=id, empresa=empresa_actual)
        trabajadores = []
        totales = {'dias': 0, 'total_general': 0, 'retiro': 0, 'patronal_rcv': 0, 'obrera_rcv': 0, 'total_rcv': 0, 'ap_pat_inf': 0, 'tipo_val_inf': 0, 'amortiz': 0, 'total_inf': 0, 'cuota_fija': 0, 'exc_pat': 0, 'exc_obr': 0, 'pd_pat': 0, 'pd_obr': 0, 'gm_pat': 0, 'gm_obr': 0, 'rt': 0, 'iv_pat': 0, 'iv_obr': 0, 'gps': 0, 'imss_pat': 0, 'imss_obr': 0, 'imss_sub': 0}
        for t in imp.trabajadores.all().order_by('id'):
            t_dict = {'nss': t.nss, 'nombre': t.nombre, 'rfc': t.rfc_curp, 'clave_u': t.clave_ubicacion, 'clave_mov': t.clave_mov, 'fecha_mov': t.fecha_mov, 'dias': t.dias, 'sdi': str(t.sdi), 'lic': t.licencias, 'inc': t.incapacidades, 'aus': t.ausentismos, 'total_general': str(t.total_general)}
            if imp.tipo == 'mensual':
                t_dict.update({'cf': str(t.cuota_fija), 'exc_pat': str(t.excedente_patronal), 'exc_obr': str(t.excedente_obrera), 'pd_pat': str(t.prestaciones_dinero_patronal), 'pd_obr': str(t.prestaciones_dinero_obrera), 'gm_pat': str(t.gastos_medicos_patronal), 'gm_obr': str(t.gastos_medicos_obrera), 'rt': str(t.riesgo_trabajo_cuota), 'iv_pat': str(t.invalidez_vida_patronal), 'iv_obr': str(t.invalidez_vida_obrera), 'gps': str(t.guarderias_ps), 'imss_pat': str(t.imss_patronal), 'imss_obr': str(t.imss_obrera), 'imss_sub': str(t.imss_subtotal)})
                totales['cuota_fija'] += float(t.cuota_fija); totales['exc_pat'] += float(t.excedente_patronal); totales['exc_obr'] += float(t.excedente_obrera); totales['pd_pat'] += float(t.prestaciones_dinero_patronal); totales['pd_obr'] += float(t.prestaciones_dinero_obrera); totales['gm_pat'] += float(t.gastos_medicos_patronal); totales['gm_obr'] += float(t.gastos_medicos_obrera); totales['rt'] += float(t.riesgo_trabajo_cuota); totales['iv_pat'] += float(t.invalidez_vida_patronal); totales['iv_obr'] += float(t.invalidez_vida_obrera); totales['gps'] += float(t.guarderias_ps); totales['imss_pat'] += float(t.imss_patronal); totales['imss_obr'] += float(t.imss_obrera); totales['imss_sub'] += float(t.imss_subtotal)
            else:
                t_dict.update({'retiro': str(t.retiro), 'patronal_rcv': str(t.patronal), 'obrera_rcv': str(t.obrera), 'total_rcv': str(t.subtotal), 'ap_pat_inf': str(t.aportacion_patronal), 'tipo_val_inf': t.tipo_valor_infonavit or '-', 'amortiz': str(t.amortizacion), 'total_inf': str(t.suma_infonavit), 'cred_viv': t.cred_vivienda or '', 'tipo_mov_cred': t.tipo_mov_credito or '', 'fecha_mov_cred': t.fecha_mov_credito or ''})
                totales['retiro'] += float(t.retiro); totales['patronal_rcv'] += float(t.patronal); totales['obrera_rcv'] += float(t.obrera); totales['total_rcv'] += float(t.subtotal); totales['ap_pat_inf'] += float(t.aportacion_patronal); totales['amortiz'] += float(t.amortizacion); totales['total_inf'] += float(t.suma_infonavit)
                if t.tipo_valor_infonavit and t.tipo_valor_infonavit != '-':
                    v_l = re.sub(r'[^\d.]', '', t.tipo_valor_infonavit)
                    if v_l: totales['tipo_val_inf'] += float(v_l)
            totales['dias'] += t.dias; totales['total_general'] += float(t.total_general); trabajadores.append(t_dict)
        for k in totales: totales[k] = int(totales[k]) if k == 'dias' else "{:,.2f}".format(totales[k])
        data = {'empresa': {'razon_social': imp.nombre_razon_social, 'rfc': imp.rfc_empresa, 'reg_patronal': imp.registro_patronal, 'actividad': limpiar_basura_header(imp.actividad), 'domicilio': limpiar_basura_header(imp.domicilio), 'cp': limpiar_basura_header(imp.cp), 'entidad': limpiar_basura_header(imp.entidad), 'periodo': limpiar_basura_header(imp.periodo), 'tipo': imp.get_tipo_display(), 'tipo_raw': imp.tipo}, 'trabajadores': trabajadores, 'totales': totales}
        return JsonResponse({'success': True, 'data': data})
    except ImportacionSUA.DoesNotExist: return JsonResponse({'success': False, 'error': 'No se encontró la importación.'})
    except Exception as e: return JsonResponse({'success': False, 'error': str(e)})

@login_required(login_url='/login/')
@require_POST
@require_hr_permission('sua', 'eliminar', json_response=True)
def eliminar_sua_ajax(request, id):
    empresa_actual = get_empresa_actual(request)
    try:
        imp = ImportacionSUA.objects.get(id=id, empresa=empresa_actual)
        crear_notificacion(
            empresa=empresa_actual,
            actor=request.user,
            mensaje=f'eliminó reporte SUA {imp.periodo}',
            link='/recursos-humanos/sua/',
            propietario=imp.creado_por or request.user
        )
        imp.delete()
        return JsonResponse({'success': True, 'message': 'Registro eliminado correctamente.'})
    except ImportacionSUA.DoesNotExist: return JsonResponse({'success': False, 'error': 'No se encontró la importación.'})
    except Exception as e: return JsonResponse({'success': False, 'error': str(e)})

@login_required(login_url='/login/')
@require_hr_permission('sua', 'exportar_excel')
def exportar_sua_excel(request, id):
    empresa_actual = get_empresa_actual(request)
    formato = request.GET.get('formato', 'excel').lower()
    try:
        imp = ImportacionSUA.objects.get(id=id, empresa=empresa_actual)
        
        meta_rows = [
            ['REPORTE DE INTEGRACIÓN SUA'],
            ['Empresa', imp.nombre_razon_social],
            ['Registro Patronal', imp.registro_patronal],
            ['Periodo', imp.periodo],
            ['Tipo', imp.get_tipo_display()],
            []
        ]
        
        if imp.tipo == 'mensual': 
            headers = ['NSS', 'Nombre', 'RFC/CURP', 'Ubicación', 'Movimiento', 'Fecha Mov.', 'Días', 'SDI', 'Lic.', 'Inc.', 'Aus.', 'C.F.', 'Exc. Pat.', 'Exc. Obr.', 'P.D. Pat.', 'P.D. Obr.', 'G.M.P. Pat.', 'G.M.P. Obr.', 'R.T.', 'I.V. Pat.', 'I.V. Obr.', 'G.P.S.', 'Patronal', 'Obrera', 'Subtotal']
        else: 
            headers = ['NSS', 'Nombre', 'RFC/CURP', 'Ubicación', 'Movimiento', 'Fecha Mov.', 'Días', 'SDI', 'Lic.', 'Inc.', 'Aus.', 'Retiro', 'Patronal RCV', 'Obrera RCV', 'Suma RCV', 'Ap. Pat. Infonavit', '%/$ /FD', 'Amortización', 'Suma Infonavit', 'Total General', 'Créd. Vivienda', 'Tipo Mov. Crédito', 'Fecha Mov. Crédito']
            
        data_rows = []
        for t in imp.trabajadores.all().order_by('id'):
            if imp.tipo == 'mensual': 
                data_rows.append([t.nss, t.nombre, t.rfc_curp, t.clave_ubicacion, t.clave_mov, t.fecha_mov, t.dias, t.sdi, t.licencias, t.incapacidades, t.ausentismos, t.cuota_fija, t.excedente_patronal, t.excedente_obrera, t.prestaciones_dinero_patronal, t.prestaciones_dinero_obrera, t.gastos_medicos_patronal, t.gastos_medicos_obrera, t.riesgo_trabajo_cuota, t.invalidez_vida_patronal, t.invalidez_vida_obrera, t.guarderias_ps, t.imss_patronal, t.imss_obrera, t.imss_subtotal])
            else: 
                data_rows.append([t.nss, t.nombre, t.rfc_curp, t.clave_ubicacion, t.clave_mov, t.fecha_mov, t.dias, t.sdi, t.licencias, t.incapacidades, t.ausentismos, t.retiro, t.patronal, t.obrera, t.subtotal, t.aportacion_patronal, t.tipo_valor_infonavit, t.amortizacion, t.suma_infonavit, t.total_general, t.cred_vivienda, t.tipo_mov_credito, t.fecha_mov_credito])

        if formato == 'csv':
            response = HttpResponse(content_type='text/csv')
            response['Content-Disposition'] = f'attachment; filename="SUA_{imp.periodo}_{imp.registro_patronal}.csv"'
            response.write(u'\ufeff'.encode('utf8'))
            writer = csv.writer(response)
            for row in meta_rows:
                writer.writerow(row)
            writer.writerow(headers)
            writer.writerows(data_rows)
            return response
        else:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "SUA Reporte"
            
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            title_font = Font(name='Calibri', size=14, bold=True, color='000000')
            meta_font = Font(name='Calibri', size=11, bold=True)
            header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
            data_font = Font(name='Calibri', size=11)
            
            header_fill = PatternFill(start_color='00b8b9', end_color='00b8b9', fill_type='solid')
            meta_fill = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
            
            thin_border = Border(
                left=Side(style='thin', color='D3D3D3'),
                right=Side(style='thin', color='D3D3D3'),
                top=Side(style='thin', color='D3D3D3'),
                bottom=Side(style='thin', color='D3D3D3')
            )
            
            ws.append(['REPORTE DE INTEGRACIÓN SUA'])
            ws.cell(row=1, column=1).font = title_font
            
            ws.append(['Empresa', imp.nombre_razon_social])
            ws.append(['Registro Patronal', imp.registro_patronal])
            ws.append(['Periodo', imp.periodo])
            ws.append(['Tipo', imp.get_tipo_display()])
            ws.append([])
            
            for r in range(2, 6):
                ws.cell(row=r, column=1).font = meta_font
                ws.cell(row=r, column=1).fill = meta_fill
                ws.cell(row=r, column=2).font = data_font
            
            ws.append(headers)
            header_row_idx = 7
            for c in range(1, len(headers) + 1):
                cell = ws.cell(row=header_row_idx, column=c)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.border = thin_border
            
            current_row = 8
            for row_data in data_rows:
                ws.append(row_data)
                for c in range(1, len(headers) + 1):
                    cell = ws.cell(row=current_row, column=c)
                    cell.font = data_font
                    cell.border = thin_border
                    if imp.tipo == 'mensual':
                        if c in [7, 8, 12, 13, 14, 15, 16, 17, 18, 19, 20, 21, 22, 23, 24, 25]:
                            cell.number_format = '$#,##0.00'
                    else:
                        if c in [7, 8, 12, 13, 14, 15, 18, 19, 20]:
                            cell.number_format = '$#,##0.00'
                current_row += 1
                
            from openpyxl.utils import get_column_letter
            for col in ws.columns:
                max_len = 0
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
                
            response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            response['Content-Disposition'] = f'attachment; filename="SUA_{imp.periodo}_{imp.registro_patronal}.xlsx"'
            wb.save(response)
            return response
            
    except ImportacionSUA.DoesNotExist: 
        return HttpResponse("No se encontró la importación", status=404)

def obtener_rango_fechas_periodo(periodo_str, tipo_sua):
    import re
    import datetime
    
    p_upper = (periodo_str or "").upper().strip()
    
    # Extraer el año
    m_anio = re.search(r'\b(20\d{2})\b', p_upper)
    anio = int(m_anio.group(1)) if m_anio else datetime.date.today().year
    
    meses_es = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", 
                "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]
                
    mes_detectado = None
    for idx, mes_nom in enumerate(meses_es, start=1):
        if mes_nom in p_upper:
            mes_detectado = idx
            break
            
    # Si no se detectó por nombre, buscar si hay algún número de 1 o 2 dígitos que sea un mes válido (1 al 12)
    if not mes_detectado:
        nums = [int(n) for n in re.findall(r'\b(\d{1,2})\b', p_upper)]
        valid_months = [n for n in nums if 1 <= n <= 12]
        if valid_months:
            mes_detectado = valid_months[0]
            
    # Caso mensual
    if tipo_sua == 'mensual' and mes_detectado:
        start_date = datetime.date(anio, mes_detectado, 1)
        if mes_detectado == 12:
            end_date = datetime.date(anio + 1, 1, 1) - datetime.timedelta(days=1)
        else:
            end_date = datetime.date(anio, mes_detectado + 1, 1) - datetime.timedelta(days=1)
        return start_date, end_date
        
    # Caso bimestral
    # Si detectamos un mes, podemos deducir el bimestre (Ene/Feb=1, Mar/Abr=2, May/Jun=3, Jul/Ago=4, Sep/Oct=5, Nov/Dic=6)
    if mes_detectado:
        bim = (mes_detectado - 1) // 2 + 1
    else:
        m_bim = re.search(r'BIMESTRE\s*(\d)', p_upper)
        if m_bim:
            bim = int(m_bim.group(1))
        else:
            m_num = re.search(r'\b([1-6])\b', p_upper)
            bim = int(m_num.group(1)) if m_num else 1
        
    if 1 <= bim <= 6:
        start_month = (bim - 1) * 2 + 1
        end_month = bim * 2
        start_date = datetime.date(anio, start_month, 1)
        if end_month == 12:
            end_date = datetime.date(anio + 1, 1, 1) - datetime.timedelta(days=1)
        else:
            end_date = datetime.date(anio, end_month + 1, 1) - datetime.timedelta(days=1)
        return start_date, end_date
        
    if mes_detectado:
        start_date = datetime.date(anio, mes_detectado, 1)
        if mes_detectado == 12:
            end_date = datetime.date(anio + 1, 1, 1) - datetime.timedelta(days=1)
        else:
            end_date = datetime.date(anio, mes_detectado + 1, 1) - datetime.timedelta(days=1)
        return start_date, end_date
        
    return datetime.date(anio, 1, 1), datetime.date(anio, 12, 31)

@login_required(login_url='/login/')
@require_POST
@require_hr_permission('sua', 'alta_empleados', json_response=True)
def alta_empleados_sua_ajax(request, id):
    empresa_actual = get_empresa_actual(request)
    try:
        importacion = ImportacionSUA.objects.get(id=id, empresa=empresa_actual)
        trabajadores = importacion.trabajadores.all(); sucursal_id = request.session.get('sucursal_id')
        creados = 0; actualizados = 0; vinculados_a_contrato = 0
        beneficiarios_map = {b.clave.strip().upper(): b for b in Beneficiario.objects.filter(empresa=empresa_actual) if b.clave}
        
        # Calcular el rango de fechas de la cédula SUA
        sua_start, sua_end = obtener_rango_fechas_periodo(importacion.periodo, importacion.tipo)

        # Buscar todos los contratos de la empresa
        contratos_candidatos = list(Contrato.objects.filter(empresa=empresa_actual))

        rfc_reporte = re.sub(r'[^A-Z0-9]', '', (importacion.rfc_empresa or '').upper()).strip()[:13]
        rp_reporte = re.sub(r'[^A-Z0-9]', '', (importacion.registro_patronal or '').upper()).strip()
        nombre_reporte = (importacion.nombre_razon_social or '').strip()

        filtros_or = Q()
        if rfc_reporte and rfc_reporte != "POR_DEFINIR": filtros_or |= Q(rfc__iexact=rfc_reporte)
        if rp_reporte: filtros_or |= Q(registro_patronal__iexact=rp_reporte) | Q(registros_patronales_adicionales__registro_patronal__iexact=rp_reporte)
        if nombre_reporte: filtros_or |= Q(nombre_razon_social__icontains=nombre_reporte)
            
        contratista_obj = Contratista.objects.filter(Q(empresa=empresa_actual) & filtros_or).distinct().first()
        status_cont = "Existente"
        if not contratista_obj:
            contratista_obj = Contratista.objects.create(empresa=empresa_actual, sucursal_id=sucursal_id, registro_patronal=rp_reporte, nombre_razon_social=nombre_reporte, rfc=rfc_reporte or "POR_DEFINIR", calle=importacion.domicilio, cp=importacion.cp, entidad_federativa=importacion.entidad, correo=f"contacto@{rp_reporte or 'empresa'}.com", creado_por=request.user)
            status_cont = "NUEVO REGISTRO"
        else:
            save_needed = False
            if (not contratista_obj.rfc or contratista_obj.rfc == "POR_DEFINIR") and rfc_reporte: contratista_obj.rfc = rfc_reporte; save_needed = True
            if not contratista_obj.registro_patronal and rp_reporte: contratista_obj.registro_patronal = rp_reporte; save_needed = True
            if save_needed: contratista_obj.save()

        with transaction.atomic():
            for t in trabajadores:
                nss_clean = re.sub(r'[^0-9]', '', t.nss).strip()[:11]
                curp_clean = re.sub(r'[^A-Z0-9]', '', (t.rfc_curp or '').upper()).strip()[:18]
                empleado = Empleado.objects.filter(Q(nss=nss_clean) | Q(curp=curp_clean), empresa=empresa_actual).first()
                beneficiario_obj = beneficiarios_map.get((t.clave_ubicacion or "").strip().upper())
                nombre_partes = t.nombre.strip().split(' '); paterno = ""; materno = ""; nombres = t.nombre
                if len(nombre_partes) >= 3: paterno = nombre_partes[0]; materno = nombre_partes[1]; nombres = " ".join(nombre_partes[2:])
                elif len(nombre_partes) == 2: paterno = nombre_partes[0]; nombres = nombre_partes[1]

                if not empleado:
                    audit_nota = f"Importado el día {importacion.fecha_importacion.strftime('%d/%m/%Y')} de la cédula {importacion.periodo} del contratista {importacion.nombre_razon_social}"
                    empleado = Empleado(empresa=empresa_actual, sucursal_id=sucursal_id, nss=nss_clean, curp=curp_clean, nombre=nombres, apellido_paterno=paterno, apellido_materno=materno, sdi=t.sdi, contratista=contratista_obj, beneficiario=beneficiario_obj, puesto="", departamento="General", clave_ubicacion=t.clave_ubicacion, notas=audit_nota, estado='activo', creado_por=request.user)
                    empleado.save(); creados += 1
                else:
                    empleado.sdi = t.sdi
                    if beneficiario_obj: empleado.beneficiario = beneficiario_obj
                    if contratista_obj: empleado.contratista = contratista_obj
                    if t.clave_ubicacion: empleado.clave_ubicacion = t.clave_ubicacion
                    empleado.save(); actualizados += 1
                
                # --- VINCULACIÓN AUTOMÁTICA CON CONTRATO DE ACUERDO AL PERIODO ---
                if beneficiario_obj:
                    contrato_match = None
                    # 1. Buscar si ya existe un contrato que coincida y cubra el periodo de la cédula
                    for c_v in contratos_candidatos:
                        if c_v.beneficiario_id == beneficiario_obj.id:
                            match_contratista = (
                                c_v.contratista_id == contratista_obj.id or
                                (c_v.contratista and (c_v.contratista.rfc == rfc_reporte or c_v.contratista.registro_patronal == rp_reporte))
                            )
                            if match_contratista:
                                start_ok = c_v.fecha_inicio <= sua_end
                                end_ok = (c_v.fecha_fin is None or c_v.fecha_fin >= sua_start)
                                if start_ok and end_ok:
                                    contrato_match = c_v
                                    break

                    # 2. Si no existe contrato que cubra este periodo, buscar contrato previo para crear la versión consecutiva
                    if not contrato_match:
                        from ..models import calcular_inicio_cuatrimestre, calcular_fin_cuatrimestre
                        cuat_start = calcular_inicio_cuatrimestre(sua_start)
                        cuat_end = calcular_fin_cuatrimestre(sua_start)
                        
                        contratos_previos = [
                            c_v for c_v in contratos_candidatos
                            if c_v.beneficiario_id == beneficiario_obj.id and (
                                c_v.contratista_id == contratista_obj.id or
                                (c_v.contratista and (c_v.contratista.rfc == rfc_reporte or c_v.contratista.registro_patronal == rp_reporte))
                            )
                        ]
                        if contratos_previos:
                            contratos_previos.sort(key=lambda x: x.fecha_inicio or sua_start, reverse=True)
                            contrato_base = contratos_previos[0]
                            contrato_match = contrato_base.crear_siguiente_version(
                                fecha_inicio=cuat_start,
                                fecha_fin=cuat_end,
                                user=request.user
                            )
                            contratos_candidatos.append(contrato_match)
                        else:
                            folio_sug = f"CONT-{(beneficiario_obj.clave or str(beneficiario_obj.id)).strip().upper()}-{cuat_start.year}-{cuat_start.month:02d}"
                            vig_def = datetime.date(cuat_start.year, 12, 31)
                            contrato_match = Contrato.objects.create(
                                empresa=empresa_actual,
                                sucursal_id=sucursal_id,
                                contratista=contratista_obj,
                                beneficiario=beneficiario_obj,
                                folio=folio_sug,
                                fecha_inicio=cuat_start,
                                fecha_fin=cuat_end,
                                vigencia_contrato=vig_def,
                                objeto_contrato=f"Servicios especializados según Cédula SUA {importacion.periodo}",
                                tipo_contrato='01',
                                creado_por=request.user
                            )
                            contratos_candidatos.append(contrato_match)

                    if contrato_match:
                        if not contrato_match.empleados.filter(id=empleado.id).exists():
                            contrato_match.empleados.add(empleado)
                            vinculados_a_contrato += 1

        crear_notificacion(
            empresa=empresa_actual,
            actor=request.user,
            mensaje='realizó alta masiva de empleados desde SUA',
            link='/recursos-humanos/empleados/',
            propietario=request.user
        )
        return JsonResponse({'success': True, 'message': f'Proceso completado: {creados} nuevos, {actualizados} actualizados. {vinculados_a_contrato} vinculaciones a contratos realizadas. Contratista: {contratista_obj.nombre_razon_social}'})
    except Exception as e: return JsonResponse({'success': False, 'error': str(e)})


@login_required(login_url='/login/')
@require_POST
@require_hr_permission('sua', 'alta_empleados', json_response=True)
def alta_empleados_cargador_sua_ajax(request, id):
    empresa_actual = get_empresa_actual(request)
    try:
        importacion = get_object_or_404(ImportacionSUA, id=id, empresa=empresa_actual)
        sucursal_id = request.session.get('sucursal_id')
        archivo = request.FILES.get('archivo')
        if not archivo:
            return JsonResponse({'success': False, 'error': 'Por favor selecciona un archivo (.xlsx o .csv).'})
        
        filename = archivo.name.lower()
        if not (filename.endswith('.xlsx') or filename.endswith('.csv')):
            return JsonResponse({'success': False, 'error': 'El archivo debe ser de formato Excel (.xlsx) o CSV (.csv).'})
        
        # 1. Identificar al Contratista desde la Cédula SUA
        sua_start, sua_end = obtener_rango_fechas_periodo(importacion.periodo, importacion.tipo)
        rfc_reporte = re.sub(r'[^A-Z0-9]', '', (importacion.rfc_empresa or '').upper()).strip()[:13]
        rp_reporte = re.sub(r'[^A-Z0-9]', '', (importacion.registro_patronal or '').upper()).strip()
        nombre_reporte = (importacion.nombre_razon_social or '').strip()

        filtros_or = Q()
        if rfc_reporte and rfc_reporte != "POR_DEFINIR":
            filtros_or |= Q(rfc__iexact=rfc_reporte)
        if rp_reporte:
            filtros_or |= Q(registro_patronal__iexact=rp_reporte) | Q(registros_patronales_adicionales__registro_patronal__iexact=rp_reporte)
        if nombre_reporte:
            filtros_or |= Q(nombre_razon_social__icontains=nombre_reporte)
            
        contratista_obj = Contratista.objects.filter(Q(empresa=empresa_actual) & filtros_or).distinct().first()
        if not contratista_obj:
            contratista_obj = Contratista.objects.create(
                empresa=empresa_actual,
                sucursal_id=sucursal_id,
                registro_patronal=rp_reporte,
                nombre_razon_social=nombre_reporte,
                rfc=rfc_reporte or "POR_DEFINIR",
                calle=importacion.domicilio,
                cp=importacion.cp,
                entidad_federativa=importacion.entidad,
                correo=f"contacto@{rp_reporte or 'empresa'}.com",
                creado_por=request.user
            )
        else:
            save_needed = False
            if (not contratista_obj.rfc or contratista_obj.rfc == "POR_DEFINIR") and rfc_reporte:
                contratista_obj.rfc = rfc_reporte
                save_needed = True
            if not contratista_obj.registro_patronal and rp_reporte:
                contratista_obj.registro_patronal = rp_reporte
                save_needed = True
            if save_needed:
                contratista_obj.save()

        # 2. Leer filas del archivo
        rows = []
        if filename.endswith('.xlsx'):
            wb = openpyxl.load_workbook(archivo, data_only=True)
            sheet = wb.active
            rows = list(sheet.iter_rows(values_only=True))
        else:
            raw_content = archivo.read()
            decoded_content = None
            for encoding in ['utf-8-sig', 'utf-8', 'latin-1', 'cp1252']:
                try:
                    decoded_content = raw_content.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
            if not decoded_content:
                decoded_content = raw_content.decode('utf-8', errors='ignore')
            reader = csv.reader(decoded_content.splitlines())
            rows = [r for r in reader if any(str(field).strip() for field in r)]

        if not rows or len(rows) < 2:
            return JsonResponse({'success': False, 'error': 'El archivo no contiene filas de datos para procesar.'})

        # Normalizar encabezados
        raw_headers = [str(h or '').strip() for h in rows[0]]
        
        def norm_header(h):
            h_clean = h.upper()
            h_clean = re.sub(r'[ÁÀÄÂ]', 'A', h_clean)
            h_clean = re.sub(r'[ÉÈËÊ]', 'E', h_clean)
            h_clean = re.sub(r'[ÍÌÏÎ]', 'I', h_clean)
            h_clean = re.sub(r'[ÓÒÖÔ]', 'O', h_clean)
            h_clean = re.sub(r'[ÚÙÜÛ]', 'U', h_clean)
            h_clean = re.sub(r'[^A-Z0-9]', '', h_clean)
            return h_clean

        norm_headers = [norm_header(h) for h in raw_headers]
        
        nss_idx = None
        nombre_idx = None
        fecha_idx = None
        ben_nombre_idx = None
        ben_clave_idx = None

        for idx, nh in enumerate(norm_headers):
            if 'NSS' in nh or 'SEGURIDADSOCIAL' in nh or 'SEGUROSOCIAL' in nh:
                if nss_idx is None: nss_idx = idx
            elif 'TRABAJADOR' in nh or 'EMPLEADO' in nh or 'NOMBRE' in nh:
                if 'BENEFICIARIO' not in nh:
                    if nombre_idx is None: nombre_idx = idx
            elif 'FECHA' in nh or 'PERIODO' in nh:
                if fecha_idx is None: fecha_idx = idx
            elif 'BENEFICIARIO' in nh or 'CLIENTE' in nh:
                if 'CLAVE' in nh or 'CODIGO' in nh:
                    if ben_clave_idx is None: ben_clave_idx = idx
                else:
                    if ben_nombre_idx is None: ben_nombre_idx = idx
            elif 'CLAVE' in nh or 'UBICACION' in nh:
                if ben_clave_idx is None: ben_clave_idx = idx

        # Positional fallback if columns are in expected order: NSS, Nombre, Fecha, Nombre Ben, Clave Ben
        if nss_idx is None and len(norm_headers) >= 1: nss_idx = 0
        if nombre_idx is None and len(norm_headers) >= 2: nombre_idx = 1
        if fecha_idx is None and len(norm_headers) >= 3: fecha_idx = 2
        if ben_nombre_idx is None and len(norm_headers) >= 4: ben_nombre_idx = 3
        if ben_clave_idx is None and len(norm_headers) >= 5: ben_clave_idx = 4

        # Pre-cargar trabajadores SUA de esta importación indexados por NSS limpio
        sua_trabajadores_map = {}
        for ts in importacion.trabajadores.all():
            nss_ts_clean = re.sub(r'[^0-9]', '', ts.nss or '').strip()[:11]
            if nss_ts_clean:
                sua_trabajadores_map[nss_ts_clean] = ts

        # Pre-cargar beneficiarios existentes de la empresa
        beneficiarios_cache = {
            b.clave.strip().upper(): b for b in Beneficiario.objects.filter(empresa=empresa_actual) if b.clave
        }

        # Pre-cargar contratos existentes
        contratos_cache = list(Contrato.objects.filter(empresa=empresa_actual, contratista=contratista_obj))

        empleados_creados = 0
        empleados_actualizados = 0
        beneficiarios_creados = 0
        contratos_creados = 0
        vinculaciones_realizadas = 0

        import datetime

        def parse_fecha_cargador(val):
            if not val:
                return None
            if isinstance(val, (datetime.date, datetime.datetime)):
                return val.date() if isinstance(val, datetime.datetime) else val
            val_str = str(val).strip()
            for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y', '%Y/%m/%d', '%d/%m/%y', '%d-%m-%y'):
                try:
                    return datetime.datetime.strptime(val_str, fmt).date()
                except ValueError:
                    continue
            return None

        with transaction.atomic():
            for row in rows[1:]:
                if not row or not any(str(c).strip() for c in row if c is not None):
                    continue

                raw_nss = str(row[nss_idx] or '') if nss_idx is not None and nss_idx < len(row) else ''
                nss_clean = re.sub(r'[^0-9]', '', raw_nss).strip()[:11]
                if not nss_clean:
                    continue

                raw_nombre = str(row[nombre_idx] or '') if nombre_idx is not None and nombre_idx < len(row) else ''
                raw_fecha = row[fecha_idx] if fecha_idx is not None and fecha_idx < len(row) else None
                raw_ben_nombre = str(row[ben_nombre_idx] or '') if ben_nombre_idx is not None and ben_nombre_idx < len(row) else ''
                raw_ben_clave = str(row[ben_clave_idx] or '') if ben_clave_idx is not None and ben_clave_idx < len(row) else ''

                clave_ben_clean = raw_ben_clave.strip().upper()
                nombre_ben_clean = raw_ben_nombre.strip()

                # Si no hay clave de beneficiario, intentar usar el nombre o marcar GENERAL
                if not clave_ben_clean and nombre_ben_clean:
                    clave_ben_clean = re.sub(r'[^A-Z0-9]', '', nombre_ben_clean.upper())[:10]
                if not clave_ben_clean:
                    clave_ben_clean = "GENERAL"

                if not nombre_ben_clean:
                    nombre_ben_clean = f"Beneficiario {clave_ben_clean}"

                # 3. Obtener o crear Beneficiario
                beneficiario_obj = beneficiarios_cache.get(clave_ben_clean)
                if not beneficiario_obj:
                    beneficiario_obj = Beneficiario.objects.filter(empresa=empresa_actual, clave__iexact=clave_ben_clean).first()
                if not beneficiario_obj:
                    beneficiario_obj = Beneficiario.objects.create(
                        empresa=empresa_actual,
                        sucursal_id=sucursal_id,
                        clave=clave_ben_clean,
                        nombre_razon_social=nombre_ben_clean,
                        creado_por=request.user
                    )
                    beneficiarios_cache[clave_ben_clean] = beneficiario_obj
                    beneficiarios_creados += 1

                # 4. Datos del trabajador desde Cédula SUA (SDI, CURP)
                ts = sua_trabajadores_map.get(nss_clean)
                curp_clean = ""
                if ts and ts.rfc_curp:
                    curp_clean = re.sub(r'[^A-Z0-9]', '', ts.rfc_curp.upper()).strip()[:18]
                sdi_val = ts.sdi if ts and ts.sdi else Decimal('0')

                # Si el nombre viene en el cargador, lo usamos; si no, el del SUA
                nombre_trabajador_final = raw_nombre.strip() or (ts.nombre.strip() if ts else f"Empleado {nss_clean}")

                nombre_partes = nombre_trabajador_final.split(' ')
                paterno = ""; materno = ""; nombres = nombre_trabajador_final
                if len(nombre_partes) >= 3:
                    paterno = nombre_partes[0]; materno = nombre_partes[1]; nombres = " ".join(nombre_partes[2:])
                elif len(nombre_partes) == 2:
                    paterno = nombre_partes[0]; nombres = nombre_partes[1]

                # 5. Crear o actualizar Empleado
                filtros_emp = Q(nss=nss_clean)
                if curp_clean:
                    filtros_emp |= Q(curp=curp_clean)
                empleado = Empleado.objects.filter(filtros_emp, empresa=empresa_actual).first()

                audit_nota = f"Alta vía Cargador Cédula SUA {importacion.periodo} ({contratista_obj.nombre_razon_social})"
                if not empleado:
                    empleado = Empleado.objects.create(
                        empresa=empresa_actual,
                        sucursal_id=sucursal_id,
                        nss=nss_clean,
                        curp=curp_clean,
                        nombre=nombres,
                        apellido_paterno=paterno,
                        apellido_materno=materno,
                        sdi=sdi_val,
                        contratista=contratista_obj,
                        beneficiario=beneficiario_obj,
                        clave_ubicacion=clave_ben_clean,
                        puesto="",
                        departamento="General",
                        notas=audit_nota,
                        estado='activo',
                        creado_por=request.user
                    )
                    empleados_creados += 1
                else:
                    empleado.contratista = contratista_obj
                    empleado.beneficiario = beneficiario_obj
                    empleado.clave_ubicacion = clave_ben_clean
                    if sdi_val > 0:
                        empleado.sdi = sdi_val
                    if curp_clean and not empleado.curp:
                        empleado.curp = curp_clean
                    empleado.save()
                    empleados_actualizados += 1

                # 6. Autogenerar o Vincular Contrato
                # Fechas del contrato: se calculan según los periodos cuatrimestrales cerrados
                from ..models import calcular_inicio_cuatrimestre, calcular_fin_cuatrimestre
                fecha_fila = parse_fecha_cargador(raw_fecha)
                if fecha_fila and (fecha_fila < sua_start or fecha_fila > sua_end):
                    c_start = calcular_inicio_cuatrimestre(fecha_fila)
                    c_end = calcular_fin_cuatrimestre(fecha_fila)
                else:
                    c_start = calcular_inicio_cuatrimestre(sua_start)
                    c_end = calcular_fin_cuatrimestre(sua_start)

                anio_folio = c_start.year
                mes_folio = f"{c_start.month:02d}"
                folio_sugerido = f"CONT-{clave_ben_clean}-{anio_folio}-{mes_folio}"

                # Buscar contrato que coincida con contratista, beneficiario y periodo
                contrato_match = None
                for c_cand in contratos_cache:
                    if c_cand.beneficiario_id == beneficiario_obj.id and c_cand.contratista_id == contratista_obj.id:
                        start_ok = c_cand.fecha_inicio <= c_end
                        end_ok = (c_cand.fecha_fin is None or c_cand.fecha_fin >= c_start)
                        if start_ok and end_ok:
                            contrato_match = c_cand
                            break

                if not contrato_match:
                    contratos_previos = [
                        c_cand for c_cand in contratos_cache
                        if c_cand.beneficiario_id == beneficiario_obj.id and c_cand.contratista_id == contratista_obj.id
                    ]
                    if contratos_previos:
                        contratos_previos.sort(key=lambda x: x.fecha_inicio or c_start, reverse=True)
                        contrato_base = contratos_previos[0]
                        contrato_match = contrato_base.crear_siguiente_version(
                            fecha_inicio=c_start,
                            fecha_fin=c_end,
                            user=request.user
                        )
                    else:
                        vig_def = datetime.date(c_start.year, 12, 31)
                        contrato_match = Contrato.objects.create(
                            empresa=empresa_actual,
                            sucursal_id=sucursal_id,
                            contratista=contratista_obj,
                            beneficiario=beneficiario_obj,
                            folio=folio_sugerido,
                            fecha_inicio=c_start,
                            fecha_fin=c_end,
                            vigencia_contrato=vig_def,
                            objeto_contrato=f"Servicios especializados según Cédula SUA {importacion.periodo} / Cargador",
                            tipo_contrato='01',
                            estado='vigente',
                            creado_por=request.user
                        )
                    contratos_cache.append(contrato_match)
                    contratos_creados += 1

                # Vincular empleado al contrato si aún no está asignado
                if not contrato_match.empleados.filter(id=empleado.id).exists():
                    contrato_match.empleados.add(empleado)
                    vinculaciones_realizadas += 1

        crear_notificacion(
            empresa=empresa_actual,
            actor=request.user,
            mensaje=f'procesó alta masiva vía Cargador SUA: {empleados_creados} nuevos empleados, {contratos_creados} contratos nuevos',
            link='/recursos-humanos/contratos/',
            propietario=request.user
        )

        msg = (
            f"Proceso completado exitosamente.\n\n"
            f"• Empleados: {empleados_creados} creados, {empleados_actualizados} actualizados.\n"
            f"• Beneficiarios: {beneficiarios_creados} creados automáticamente.\n"
            f"• Contratos: {contratos_creados} autogenerados.\n"
            f"• Vinculaciones a Contratos: {vinculaciones_realizadas} asignaciones.\n"
            f"• Contratista: {contratista_obj.nombre_razon_social}"
        )
        return JsonResponse({
            'success': True,
            'message': msg,
            'stats': {
                'empleados_creados': empleados_creados,
                'empleados_actualizados': empleados_actualizados,
                'beneficiarios_creados': beneficiarios_creados,
                'contratos_creados': contratos_creados,
                'vinculaciones_realizadas': vinculaciones_realizadas,
                'contratista': contratista_obj.nombre_razon_social
            }
        })
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)})


@login_required(login_url='/login/')
def descargar_plantilla_cargador_sua(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cargador_Empleados_SUA"
    
    headers = [
        "NSS", 
        "Nombre del Trabajador", 
        "Fecha de Trabajo Realizado", 
        "Nombre Beneficiario", 
        "Clave Beneficiario"
    ]
    
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    
    fill_header = PatternFill(start_color="00B8B9", end_color="00B8B9", fill_type="solid")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center")
    border_thin = Border(
        left=Side(style='thin', color="CCCCCC"),
        right=Side(style='thin', color="CCCCCC"),
        top=Side(style='thin', color="CCCCCC"),
        bottom=Side(style='thin', color="CCCCCC")
    )
    
    ws.row_dimensions[1].height = 26
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = align_center
        cell.border = border_thin
        
    sample_data = [
        ["12345678901", "HERNANDEZ LOPEZ CARLOS", "01/01/2026", "SERVICIOS INDUSTRIALES ADT", "ADT01"],
        ["98765432109", "MARTINEZ GARCIA SOFIA", "01/01/2026", "SERVICIOS INDUSTRIALES ADT", "ADT01"],
    ]
    
    for row_idx, row_data in enumerate(sample_data, start=2):
        ws.row_dimensions[row_idx].height = 20
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = border_thin
            if col_idx in [1, 3, 5]:
                cell.alignment = align_center
                
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 35
    ws.column_dimensions['C'].width = 28
    ws.column_dimensions['D'].width = 35
    ws.column_dimensions['E'].width = 22
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="Plantilla_Cargador_SUA_Empleados.xlsx"'
    wb.save(response)
    return response

