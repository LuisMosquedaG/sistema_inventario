import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = docx.Document()

# Page setup
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

# Brand Colors
COLOR_BRAND = RGBColor(0, 184, 185)       # #00B8B9 (Teal)
COLOR_DARK = RGBColor(40, 45, 50)         # #282D32
COLOR_MUTED = RGBColor(108, 117, 125)     # #6C757D
COLOR_SUCCESS = RGBColor(47, 158, 68)     # #2F9E44
COLOR_WARNING = RGBColor(245, 159, 0)     # #F59F00
COLOR_DANGER = RGBColor(224, 49, 49)      # #E03131

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def add_header_title(doc, text, subtitle=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.space_before = Pt(10)
    run = p.add_run(text)
    run.font.name = 'Segoe UI'
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = COLOR_BRAND

    if subtitle:
        p_sub = doc.add_paragraph()
        p_sub.paragraph_format.space_after = Pt(14)
        run_sub = p_sub.add_run(subtitle)
        run_sub.font.name = 'Segoe UI'
        run_sub.font.size = Pt(11)
        run_sub.font.italic = True
        run_sub.font.color.rgb = COLOR_MUTED

def add_section_h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = 'Segoe UI'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = COLOR_DARK

def add_paragraph(doc, text, bold_prefix=None, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Segoe UI'
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_DARK
    r_body = p.add_run(text)
    r_body.font.name = 'Segoe UI'
    r_body.font.size = Pt(10)
    r_body.font.color.rgb = COLOR_DARK
    return p

def add_callout(doc, text, title="NOTA"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.8)
    set_cell_background(cell, "F1F9F9")
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    
    r_title = p.add_run(f"📌 {title}: ")
    r_title.font.name = 'Segoe UI'
    r_title.font.size = Pt(9.5)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_BRAND
    
    r_text = p.add_run(text)
    r_text.font.name = 'Segoe UI'
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = COLOR_DARK
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_screenshot_box(doc, caption):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.8)
    set_cell_background(cell, "F8F9FA")
    
    # Border styling
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="CED4DA"/>
            <w:left w:val="single" w:sz="6" w:space="0" w:color="CED4DA"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="CED4DA"/>
            <w:right w:val="single" w:sz="6" w:space="0" w:color="CED4DA"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(20)
    
    r_box = p.add_run(f"📷 ESPACIO PARA CAPTURA DE PANTALLA\n\n[ {caption} ]")
    r_box.font.name = 'Segoe UI'
    r_box.font.size = Pt(10)
    r_box.font.bold = True
    r_box.font.color.rgb = COLOR_MUTED
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

# --- DOCUMENT CONTENT ---
add_header_title(doc, "Guía de Usuario: Portal de Proveedores", "Gestión y Carga de Documentación de Cumplimiento Mensual")

add_section_h2(doc, "1. Roles y Terminología en el Sistema")
add_paragraph(doc, " La empresa que presta el servicio y que es responsable de subir oportunamente sus comprobantes y reportes periódicos.", bold_prefix="• Tú (El Proveedor):")
add_paragraph(doc, " La empresa contratante con la que colaboras, encargada de verificar, autorizar o solicitar correcciones a los documentos que subes.", bold_prefix="• Tu Cliente (El Contratista):")

add_section_h2(doc, "2. Acceso al Portal")
add_paragraph(doc, "1. Abre tu navegador web favorito (Chrome, Edge, Firefox, Safari).")
add_paragraph(doc, "2. Ingresa a la URL oficial del portal: https://suite.crossovermx.com/login/")
add_paragraph(doc, "3. Captura tus credenciales proporcionadas por tu cliente:")
add_paragraph(doc, "   • Usuario: tu_usuario@subdominio_cliente (Ej. comercial@constructora)")
add_paragraph(doc, "   • Contraseña: Tu contraseña asignada.")
add_paragraph(doc, "4. Haz clic en 'Iniciar Sesión'. El sistema te redirigirá directamente a tu Portal de Documentación.")

add_callout(doc, "Si no cuentas con tu usuario y contraseña, o si necesitas restablecerla, solicítala directamente al área de compras o administración de tu cliente.", "IMPORTANTE")
add_screenshot_box(doc, "CAPTURA 1: Pantalla de Inicio de Sesión (Login)")

add_section_h2(doc, "3. Matriz de Documentación (Expediente Anual)")
add_paragraph(doc, "Al ingresar, encontrarás la Matriz de Documentación organizada en una tabla con los siguientes elementos:")
add_paragraph(doc, " Permite consultar o cargar comprobantes de años anteriores o del año en curso.", bold_prefix="• Selector de Año Fiscal: ")
add_paragraph(doc, " Requisitos obligatorios solicitados por tu cliente (Opinión SAT, SUA, REPSE, Reporte de trabajadores, etc.).", bold_prefix="• Columna de Documentos: ")
add_paragraph(doc, " Cada celda representa el mes en el que debes adjuntar el comprobante correspondiente.", bold_prefix="• Cuadrícula de Meses (Ene - Dic): ")

add_paragraph(doc, "Periodicidad de entrega:")
add_paragraph(doc, "   - Mensual: Requiere un comprobante por cada mes.")
add_paragraph(doc, "   - Bimestral / Cuatrimestral: La celda abarca el periodo correspondiente (ej. Ene-Abr).")
add_paragraph(doc, "   - Única ocasión: Documentos fijos o de vigencia anual.")

add_screenshot_box(doc, "CAPTURA 2: Matriz de Documentos por Mes y Selector de Año")

add_section_h2(doc, "4. Carga, Consulta y Descarga de Documentos")
add_paragraph(doc, "Para cargar un archivo, haz clic sobre la celda vacía (marcada con un guion '—') del documento y mes que deseas comprobar. Selecciona el archivo de tu equipo y el sistema lo subirá automáticamente.", bold_prefix="• Cargar Documento: ")
add_paragraph(doc, "Junto al documento 'Reporte de trabajadores' encontrarás el botón verde 'Plantilla'. Haz clic para descargar el formato oficial en Excel, llénalo con los datos de tu personal y súbelo al mes correspondiente.", bold_prefix="• Descargar Plantilla Excel: ")
add_paragraph(doc, "Haz clic sobre cualquier celda con archivo cargado para abrirlo y revisarlo en una pestaña nueva.", bold_prefix="• Consultar o Descargar: ")
add_paragraph(doc, "Si necesitas sustituir un archivo antes de su aprobación, pulsa el icono de la cruz (✕) en la celda y sube el archivo correcto.", bold_prefix="• Reemplazar / Eliminar: ")

add_screenshot_box(doc, "CAPTURA 3: Carga de Archivo y Botón de Plantilla")

add_section_h2(doc, "5. Estados y Estatus de los Documentos")
add_paragraph(doc, "Cada comprobante dentro de la matriz pasa por un ciclo de validación visual:")

# Table of Status
tbl = doc.add_table(rows=5, cols=3)
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

headers = ["Estatus", "Color Visual", "Significado para el Proveedor"]
col_widths = [Inches(1.5), Inches(1.5), Inches(3.8)]

hdr_cells = tbl.rows[0].cells
for i, h in enumerate(headers):
    hdr_cells[i].text = h
    hdr_cells[i].width = col_widths[i]
    set_cell_background(hdr_cells[i], "00B8B9")
    p = hdr_cells[i].paragraphs[0]
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
    p.runs[0].font.name = 'Segoe UI'

rows_data = [
    ("Sin documento", "Gris suave", "Documento pendiente por subir para el periodo."),
    ("En revisión", "Amarillo suave", "Documento subido correctamente. Tu cliente lo está revisando."),
    ("Aprobado", "Verde suave", "Tu cliente validó el documento y cumple todos los requisitos."),
    ("Rechazado", "Rojo suave", "Tu cliente detectó inconsistencias. Revisa el motivo y vuelve a subirlo.")
]

for row_idx, data in enumerate(rows_data, start=1):
    row_cells = tbl.rows[row_idx].cells
    bg_color = "FFFFFF" if row_idx % 2 != 0 else "F8F9FA"
    for col_idx, text in enumerate(data):
        row_cells[col_idx].text = text
        row_cells[col_idx].width = col_widths[col_idx]
        set_cell_background(row_cells[col_idx], bg_color)
        p = row_cells[col_idx].paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.runs[0].font.name = 'Segoe UI'
        p.runs[0].font.size = Pt(9)

doc.add_paragraph().paragraph_format.space_after = Pt(6)
add_screenshot_box(doc, "CAPTURA 4: Estatus de Celdas y Barra de Resumen al Pie")

add_section_h2(doc, "6. Notificaciones y Correos Automáticos")
add_paragraph(doc, "El portal mantiene informadas a ambas partes mediante notificaciones automáticas por correo:")
add_paragraph(doc, "Cuando subes un documento, tu cliente recibe un correo notificándole que tiene un nuevo archivo listo para revisión.", bold_prefix="1. Aviso de Carga: ")
add_paragraph(doc, "Cuando tu cliente revisa y autoriza tu comprobante, recibes una notificación confirmando que fue aprobado.", bold_prefix="2. Aviso de Aprobación: ")
add_paragraph(doc, "Si tu cliente rechaza un archivo, recibirás un correo con el motivo y observaciones exactas para que puedas corregirlo de inmediato.", bold_prefix="3. Aviso de Rechazo con Motivo: ")
add_paragraph(doc, "Tu cliente puede enviarte recordatorios con la lista de documentos que tienes pendientes.", bold_prefix="4. Recordatorios Periódicos: ")

add_screenshot_box(doc, "CAPTURA 5: Correo Automático de Notificación Recibido")

# Save document
output_path = r"C:\Users\luism\.gemini\antigravity\brain\48938170-957f-4afe-9242-a9e967d9a78b\guia_usuario_portal_proveedores.docx"
doc.save(output_path)
print(f"Document saved successfully at: {output_path}")
